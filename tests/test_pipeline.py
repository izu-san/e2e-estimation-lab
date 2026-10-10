"""CLIを工程順に通す統合テスト。AI提出JSONは人工の固定データ。"""

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conftest import ADAPTER
from estimator.cli import main
from helpers import build_cases_standard, build_common_standard, build_customer, write_config, write_json

AX = ["AX-01", "AX-02", "AX-03", "AX-04", "AX-05"]
REF = lambda rng, sheet="テスト項目": {"workbook_id": "W1", "sheet": sheet, "range": rng}  # noqa: E731


def run_cli(capsys, *args):
    code = main(list(args))
    out = capsys.readouterr().out
    try:
        return code, json.loads(out)
    except json.JSONDecodeError:
        return code, out


def normalized():
    return {
        "schema_version": 1,
        "regions": [
            {"region_id": "W1-S1-R001", "role": "description", "disposition": "not_test", "reason": "説明文"},
            {"region_id": "W1-S1-R002", "role": "test_table", "disposition": "parsed", "header_rows": [4], "reason": "表"},
            {"region_id": "W1-S1-R003", "role": "reference", "disposition": "not_test", "reason": "参考行"},
            {"region_id": "W1-S2-R001", "role": "other", "disposition": "needs_question", "reason": "非表示シート",
             "question_ids": ["Q1"]},
        ],
        "source_items": [
            {"item_id": "T-1", "original_id": "T-1", "id_assigned": False, "region_ids": ["W1-S1-R002"],
             "refs": [REF("A5:D6")], "title": "商品詳細", "step_count": 2, "expected_count": 2,
             "data_pattern_count": 1, "disposition": "mapped", "duplicate_of": None, "reason": "",
             "inheritance": [{"ref": REF("B6"), "from_ref": REF("B5"), "reason": "結合セル"}]},
            {"item_id": "T-2", "original_id": "T-2", "id_assigned": False, "region_ids": ["W1-S1-R002"],
             "refs": [REF("A7:D7")], "title": "ログイン", "step_count": 1, "expected_count": 1,
             "data_pattern_count": 1, "disposition": "mapped", "duplicate_of": None, "reason": "非表示行だが項目"},
        ],
        "cases": [
            {"case_id": "C-001", "title": "商品詳細表示", "source_item_ids": ["T-1"], "relation": "one_to_one",
             "summary": "", "steps": ["一覧", "選択"], "expected": ["一覧", "詳細"], "step_count": 2,
             "assertion_count": 2, "data_pattern_count": 1},
            {"case_id": "C-002", "title": "ログイン画面遷移", "source_item_ids": ["T-2"], "relation": "one_to_one",
             "summary": "", "steps": ["押す"], "expected": ["画面"], "step_count": 1, "assertion_count": 1,
             "data_pattern_count": 1},
        ],
    }


def design():
    base = {"responsibility": "r", "need_reason": "n", "used_by": [], "depends_on": [], "integrated_into": None}
    return {"schema_version": 1,
            "architectures": [{"arch_id": "A1", "name": "POMあり", "assumptions": []}],
            "components": [
                {**base, "component_id": "BASE-1", "arch_ids": ["A1"], "type_id": "BASE", "name": "基盤",
                 "implementation_kind": "新規", "elements": ["BASE"], "cost_owner": "BASE-1", "used_by": ["C-001", "C-002"]},
                {**base, "component_id": "PAGE-1", "arch_ids": ["A1"], "type_id": "PAGE", "name": "ページ",
                 "implementation_kind": "新規", "elements": ["PG-List", "PG-Detail"], "cost_owner": "PAGE-1",
                 "depends_on": ["BASE-1"], "used_by": ["C-001"]},
            ]}


def assessments():
    def ax(*s):
        return {a: {"score": v, "condition": "c" if v else "", "evidence": "e" if v else ""} for a, v in zip(AX, s)}
    return {"schema_version": 1,
            "case_assessments": [
                {"case_id": "C-001", "arch_ids": ["A1"], "scope": "対象内", "scope_reason": "", "feasibility": "feasible",
                 "info_sufficient": True, "axes": ax(2, 1, 1, 1, 1), "estimate_status": "確定可能", "rationale": "r",
                 "uses_components": ["BASE-1", "PAGE-1"]},
                {"case_id": "C-002", "arch_ids": ["A1"], "scope": "対象内", "scope_reason": "", "feasibility": "feasible",
                 "info_sufficient": False, "axes": ax(1, 1, None, 1, 1), "estimate_status": "要確認",
                 "rationale": "ログイン手段不明", "question_ids": ["Q1"], "uses_components": ["BASE-1"]},
            ],
            "common_assessments": [
                {"component_id": "BASE-1", "difficulty": "Low", "matched_conditions": "m", "rationale": "r",
                 "estimate_status": "確定可能", "standard_refs": ["common:工数表!A2"]},
                {"component_id": "PAGE-1", "difficulty": "Mid", "matched_conditions": "m", "rationale": "r",
                 "estimate_status": "確定可能"},
            ]}


def questions():
    return {"schema_version": 1, "items": [
        {"question_id": "Q1", "kind": "question", "targets": ["C-002", "W1-S2-R001"], "text": "ログイン方法は？",
         "impact": "C-002が要確認", "status": "open"}]}


@pytest.fixture
def run(tmp_path, capsys):
    c = build_cases_standard(tmp_path / "cases_v0.1.xlsx")
    m = build_common_standard(tmp_path / "common_v0.1.xlsx")
    cust = build_customer(tmp_path / "cust.xlsx")
    cfg = write_config(tmp_path / "cfg.yaml", c, m, ADAPTER)
    code, out = run_cli(capsys, "init-run", "--input", str(cust), "--config", str(cfg), "--run-id", "t1")
    assert code == 0, out
    assert run_cli(capsys, "extract", "--run", "t1")[0] == 0
    rd = Path(out["run_dir"])
    write_json(rd / "ai/normalized_cases.json", normalized())
    write_json(rd / "ai/common_design.json", design())
    write_json(rd / "ai/assessments.json", assessments())
    write_json(rd / "ai/questions.json", questions())
    return rd, (c, m, cust)


def test_full_pipeline(run, capsys):
    rd, srcs = run
    from estimator.io_utils import sha256_file
    before = [sha256_file(p) for p in srcs]
    for st in ("normalize", "common_design", "questions", "assessments"):  # questionsは採点が参照するため先
        code, out = run_cli(capsys, "validate-stage", "--run", "t1", "--stage", st)
        assert code == 0, out
    code, out = run_cli(capsys, "calculate", "--run", "t1")
    assert code == 0, out
    a = out["architectures"][0]
    # C-001: 合計6 → Mid(3/5/8)。C-002: 要確認。BASE Low(1/2/3)×1 + PAGE Mid(2/4/6)×2
    assert a["subtotals"]["cases"] == {"low": 3.0, "std": 5.0, "high": 8.0}
    assert a["subtotals"]["common"] == {"low": 5.0, "std": 10.0, "high": 15.0}
    assert a["estimate_completeness"] == "一部未算定"
    code, out = run_cli(capsys, "validate", "--run", "t1")
    assert code == 0, out
    assert out["status"]["execution_status"] == "処理成功"
    assert out["status"]["semantic_review_status"] == "未実施"
    code, out = run_cli(capsys, "export", "--run", "t1")
    assert code == 0, out
    wb = load_workbook(rd / "estimate.xlsx")
    assert {"サマリー", "個別ケース", "共通作業", "原項目対応", "確認事項", "検証結果", "実行条件", "レビュー"} <= set(wb.sheetnames)
    ws = wb["個別ケース"]
    hdr = [c.value for c in ws[1]]
    row2 = {h: c.value for h, c in zip(hdr, ws[3])}  # C-002
    assert row2["case_id"] == "C-002" and row2["標準"] is None and row2["estimate_status"] == "要確認"
    assert (rd / "summary.md").read_text().startswith("# 見積もり結果")
    assert [sha256_file(p) for p in srcs] == before


def test_hand_edited_estimate_is_rejected(run, capsys):
    rd, _ = run
    for st in ("normalize", "common_design", "assessments"):
        run_cli(capsys, "validate-stage", "--run", "t1", "--stage", st)
    run_cli(capsys, "calculate", "--run", "t1")
    est = json.loads((rd / "estimate.json").read_text())
    est["architectures"][0]["subtotals"]["computed_total"]["std"] = 1.0
    (rd / "estimate.json").write_text(json.dumps(est, ensure_ascii=False))
    code, out = run_cli(capsys, "validate", "--run", "t1")
    assert code == 1
    assert any(e["code"] == "RECOMPUTE_MISMATCH" for e in out["errors"])
    assert out["status"]["execution_status"] == "処理失敗"


def test_ai_cannot_submit_computed_values(run, capsys):
    rd, _ = run
    a = assessments()
    a["case_assessments"][0]["difficulty"] = "Low"
    a["case_assessments"][0]["effort"] = {"std": 1}
    write_json(rd / "ai/assessments.json", a)
    code, out = run_cli(capsys, "validate-stage", "--run", "t1", "--stage", "assessments")
    assert code == 1 and out["errors"][0]["code"] == "SCHEMA"


def test_assessments_require_questions_first(run, capsys):
    """採点が参照する質問IDは questions.json に先に存在する必要がある（Skill工程5→6の順序）。"""
    rd, _ = run
    (rd / "ai/questions.json").unlink()
    code, out = run_cli(capsys, "validate-stage", "--run", "t1", "--stage", "assessments")
    assert code == 1
    assert {e["code"] for e in out["errors"]} == {"UNKNOWN_QUESTION"}
    write_json(rd / "ai/questions.json", questions())
    code, out = run_cli(capsys, "validate-stage", "--run", "t1", "--stage", "questions")
    assert code == 0, out
    code, out = run_cli(capsys, "validate-stage", "--run", "t1", "--stage", "assessments")
    assert code == 0, out


def test_mark_review_context(run, capsys):
    rd, _ = run
    for st in ("normalize", "common_design", "questions", "assessments"):
        assert run_cli(capsys, "validate-stage", "--run", "t1", "--stage", st)[0] == 0
    assert run_cli(capsys, "calculate", "--run", "t1")[0] == 0
    write_json(rd / "ai/semantic_review.json", {"schema_version": 1, "reviewer": "ai", "findings": []})
    code, out = run_cli(capsys, "mark-review", "--run", "t1", "--by", "ai")
    assert code != 0  # --context が必須
    code, out = run_cli(capsys, "mark-review", "--run", "t1", "--by", "ai", "--context", "self")
    assert code == 0 and out["semantic_review_status"] == "AIレビュー済（自己レビュー）"
    code, out = run_cli(capsys, "mark-review", "--run", "t1", "--by", "ai", "--context", "independent")
    assert code == 0 and out["semantic_review_status"] == "AIレビュー済"


def test_semantic_review_validation(run, capsys):
    rd, _ = run
    def review(findings):
        write_json(rd / "ai/semantic_review.json", {"schema_version": 1, "reviewer": "ai", "findings": findings})
        return run_cli(capsys, "validate-stage", "--run", "t1", "--stage", "semantic_review")
    ok = {"target": "C-001", "related": ["C-002"], "kind": "scoring", "message": "根拠が弱い", "suggestion": "再確認"}
    assert review([ok])[0] == 0
    assert review([{**ok, "target": "overall", "related": []}])[0] == 0
    code, out = review([{**ok, "target": "C-999"}])
    assert code == 1 and out["errors"][0]["code"] == "UNKNOWN_TARGET"
    code, out = review([{**ok, "related": ["PAGE-xxx"]}])
    assert code == 1 and out["errors"][0]["code"] == "UNKNOWN_TARGET"
    assert review([{**ok, "kind": "split_judgement"}])[1]["errors"][0]["code"] == "SCHEMA"  # 未定義のkind
    assert review([{**ok, "message": ""}])[1]["errors"][0]["code"] == "SCHEMA"
    # 検証エラーのままではAIレビュー済にできない
    review([{**ok, "target": "C-999"}])
    code, out = run_cli(capsys, "mark-review", "--run", "t1", "--by", "ai", "--context", "independent")
    assert code != 0


def test_warning_summary_groups_without_dropping(run, capsys):
    from estimator.validate import summarize_warnings, warning_kind
    mk = lambda st, code, msg: {"code": code, "path": "p", "message": msg, "severity": "warning"}
    machine = {"stages": {
        "standards": {"errors": [], "warnings": [
            mk("s", "STANDARDS", "読取警告 FORMULA_NO_CACHE 判定例!K9"),
            mk("s", "STANDARDS", "読取警告 FORMULA_NO_CACHE 判定例!M9"),
            mk("s", "STANDARDS", "個別工数表: 難易度 'Impossible' は対象外の行として扱います (工数標準!A5)")]},
        "assessments": {"errors": [], "warnings": [mk("s", "UNASSESSED", "C-009 の採点がありません")]}}}
    got = {(g["stage"], g["kind"]): g["count"] for g in summarize_warnings(machine)}
    assert got == {("standards", "FORMULA_NO_CACHE"): 2, ("standards", "STD_NON_EFFORT_ROW"): 1,
                   ("assessments", "UNASSESSED"): 1}
    assert sum(got.values()) == 4  # 集約しても件数は減らない（警告を捨てない）
    assert warning_kind(mk("s", "STANDARDS", "未知の警告")) == "STANDARDS"  # 未分類は元のコードのまま


def test_region_not_disposed_and_repair_limit(run, capsys):
    rd, _ = run
    n = normalized()
    n["regions"].pop()  # 非表示シートの領域を無言で捨てる
    write_json(rd / "ai/normalized_cases.json", n)
    statuses = []
    for _ in range(3):
        code, out = run_cli(capsys, "validate-stage", "--run", "t1", "--stage", "normalize")
        statuses.append(out["status"])
        assert any(e["code"] == "REGION_NOT_DISPOSED" for e in out["errors"])
    assert statuses == ["failed", "failed", "failed_repair_limit"]  # 初回 + 修正2回


def test_double_count_detected(run, capsys):
    rd, _ = run
    d = design()
    d["components"].append({**d["components"][1], "component_id": "PAGE-2", "cost_owner": "PAGE-2",
                              "elements": ["PG-List"]})
    write_json(rd / "ai/common_design.json", d)
    code, out = run_cli(capsys, "validate-stage", "--run", "t1", "--stage", "common_design")
    assert code == 1 and any(e["code"] == "DOUBLE_COUNT" for e in out["errors"])


def test_duplicate_requires_reason_and_target(run, capsys):
    rd, _ = run
    n = normalized()
    n["source_items"].append({**n["source_items"][1], "item_id": "DUP-1", "original_id": None, "id_assigned": True,
                              "disposition": "merged_duplicate", "duplicate_of": "NOPE", "reason": ""})
    write_json(rd / "ai/normalized_cases.json", n)
    code, out = run_cli(capsys, "validate-stage", "--run", "t1", "--stage", "normalize")
    codes = {e["code"] for e in out["errors"]}
    assert {"BAD_DUPLICATE", "NO_REASON"} <= codes


def test_bad_refs_detected(run, capsys):
    rd, _ = run
    n = normalized()
    n["source_items"][0]["refs"] = [REF("A500:D501")]
    n["source_items"][1]["refs"] = [REF("A1", sheet="存在しない")]
    write_json(rd / "ai/normalized_cases.json", n)
    code, out = run_cli(capsys, "validate-stage", "--run", "t1", "--stage", "normalize")
    assert sum(e["code"] == "BAD_REF" for e in out["errors"]) == 2


def test_source_change_detected_and_stale(run, capsys, tmp_path):
    rd, (c, m, cust) = run
    for st in ("normalize", "common_design", "assessments"):
        run_cli(capsys, "validate-stage", "--run", "t1", "--stage", st)
    run_cli(capsys, "calculate", "--run", "t1")
    # 共通設計を変更 → 下流（採点・計算）が無効化される
    d = design()
    d["components"][1]["elements"].append("PG-Cart")
    write_json(rd / "ai/common_design.json", d)
    code, out = run_cli(capsys, "status", "--run", "t1")
    assert out["stages"]["common_design"]["status"] == "modified"
    assert out["stages"]["assessments"]["status"] == "stale"
    assert out["stages"]["calculate"]["status"] == "stale"
    # 入力ファイルの変更を検出
    build_customer(cust, with_desc=False)
    code, out = run_cli(capsys, "extract", "--run", "t1")
    assert code == 2 and out["errors"][0]["code"] == "SOURCE_CHANGED"


def test_multiple_inputs_require_manifest(tmp_path, capsys):
    with pytest.raises(SystemExit):
        main(["init-run", "--input", "a.xlsx", "--project", "p.yaml", "--config", "c.yaml"])


def test_project_manifest(tmp_path, capsys):
    c = build_cases_standard(tmp_path / "cases_v0.1.xlsx")
    m = build_common_standard(tmp_path / "common_v0.1.xlsx")
    a = build_customer(tmp_path / "a.xlsx")
    b = build_customer(tmp_path / "b.xlsx", sheet="機能B")
    cfg = write_config(tmp_path / "cfg.yaml", c, m, ADAPTER)
    (tmp_path / "p.yaml").write_text(f"schema_version: 1\nproject_name: demo\ninputs:\n  - path: {a}\n  - path: {b}\n")
    code, out = run_cli(capsys, "init-run", "--project", str(tmp_path / "p.yaml"), "--config", str(cfg), "--run-id", "pj")
    assert code == 0
    code, out = run_cli(capsys, "extract", "--run", "pj")
    assert out["workbooks"] == 2


def test_missing_input_recorded(tmp_path, capsys):
    c = build_cases_standard(tmp_path / "cases_v0.1.xlsx")
    m = build_common_standard(tmp_path / "common_v0.1.xlsx")
    cfg = write_config(tmp_path / "cfg.yaml", c, m, ADAPTER)
    code, out = run_cli(capsys, "init-run", "--input", str(tmp_path / "none.xlsx"), "--config", str(cfg), "--run-id", "x")
    assert code == 2 and out["missing_inputs"]
    code, out = run_cli(capsys, "extract", "--run", "x")
    assert out["constraints"][0]["code"] == "NOT_FOUND"


def _stage(capsys, stage):
    code, out = run_cli(capsys, "validate-stage", "--run", "t1", "--stage", stage)
    return code, {e["code"] for e in out["errors"]}, {w["code"] for w in out.get("warnings", [])}


def test_usage_mismatch_between_design_and_assessment(run, capsys):
    """common_design の used_by と assessments の uses_components が食い違うとエラー。"""
    rd, _ = run
    a = assessments()
    a["case_assessments"][0]["uses_components"] = ["PAGE-1"]  # BASE-1 を書き漏らす
    write_json(rd / "ai/assessments.json", a)
    code, errors, _ = _stage(capsys, "assessments")
    assert code == 1 and "USAGE_MISMATCH" in errors


def test_question_links_and_bare_evidence_ref_warn(run, capsys):
    """質問のtargetsと question_ids の食い違い、シート名なしの座標は警告（エラーにはしない）。"""
    rd, _ = run
    a = assessments()
    a["case_assessments"][1]["question_ids"] = []  # Q1 の targets に C-002 があるのに参照しない
    a["case_assessments"][0]["axes"]["AX-01"]["evidence"] = "一覧と詳細（C3）"
    write_json(rd / "ai/assessments.json", a)
    code, errors, warns = _stage(capsys, "assessments")
    assert code == 0 and not errors
    assert {"QUESTION_NOT_LINKED", "EVIDENCE_REF_NO_SHEET"} <= warns


def test_step_count_mismatch_between_parent_and_children_warns(run, capsys):
    rd, _ = run
    n = normalized()
    n["source_items"][0]["step_count"] = 5
    n["cases"] = [dict(n["cases"][0], case_id=f"C-00{i}", relation="split", relation_reason="分解",
                       step_count=2) for i in (1, 3)] + [n["cases"][1]]
    for c in n["cases"][:2]:
        c["source_item_ids"] = ["T-1"]
    write_json(rd / "ai/normalized_cases.json", n)
    code, errors, warns = _stage(capsys, "normalize")
    assert code == 0 and "STEP_COUNT_MISMATCH" in warns


def test_grounding_warnings(run, capsys):
    """基準セル参照なし・推測表現・依存先の質問の引継ぎ漏れ・根拠中の未使用部品は警告（エラーにはしない）。"""
    rd, _ = run
    a = assessments()
    # C-001: 確定可能のまま、依存先PAGE-1が未解決質問に紐づく。根拠にBASE-1以外の未使用部品も書く
    a["common_assessments"][1]["question_ids"] = ["Q1"]
    a["case_assessments"][0]["axes"]["AX-01"]["evidence"] = "OTHER-1 を使う"
    # C-002: 要確認のケースで推測表現
    a["case_assessments"][1]["axes"]["AX-01"]["evidence"] = "単一操作と想定"
    d = design()
    d["components"].append({**d["components"][0], "component_id": "OTHER-1", "type_id": "UTIL", "name": "補助",
                            "elements": ["U1"], "cost_owner": "OTHER-1", "used_by": []})
    q = questions()
    q["items"][0]["targets"] = ["C-002", "PAGE-1"]
    write_json(rd / "ai/assessments.json", a)
    write_json(rd / "ai/common_design.json", d)
    write_json(rd / "ai/questions.json", q)
    code, errors, warns = _stage(capsys, "assessments")
    assert code == 0 and not errors
    assert {"NO_STANDARD_REF", "SPECULATIVE_EVIDENCE", "QUESTION_NOT_INHERITED",
            "EVIDENCE_COMPONENT_NOT_USED"} <= warns


def test_invented_numbers_are_rejected(run, capsys):
    """AIが判断文に工数・割合・合計点を書くとエラー（根拠evidenceの引用は警告）。"""
    rd, _ = run
    a = assessments()
    a["case_assessments"][0]["rationale"] = "標準16人時が妥当。共通化で30%割引"
    a["common_assessments"][0]["rationale"] = "合計8点のためHigh"
    a["case_assessments"][0]["axes"]["AX-02"]["evidence"] = "原文に「3時間以内」とある"
    write_json(rd / "ai/assessments.json", a)
    code, out = run_cli(capsys, "validate-stage", "--run", "t1", "--stage", "assessments")
    assert code == 1
    paths = {e["path"] for e in out["errors"] if e["code"] == "UNSOURCED_NUMBER"}
    assert any(p.endswith("/rationale") for p in paths) and len(paths) >= 2
    assert any(w["code"] == "UNSOURCED_NUMBER" and w["path"].endswith("/evidence") for w in out["warnings"])


def test_integrated_component_is_not_pending(run, capsys):
    """他の作業へ統合して計上した部品は『他へ計上』と表示し、要確認と誤読させない。"""
    rd, _ = run
    d = design()
    d["components"][1]["integrated_into"] = "BASE-1"
    d["components"][1]["cost_owner"] = "BASE-1"
    a = assessments()
    a["common_assessments"] = a["common_assessments"][:1]
    write_json(rd / "ai/common_design.json", d)
    write_json(rd / "ai/assessments.json", a)
    code, out = run_cli(capsys, "calculate", "--run", "t1")
    assert code == 0, out
    import json as _j
    est = _j.loads((rd / "estimate.json").read_text(encoding="utf-8"))
    row = next(r for r in est["architectures"][0]["components"] if r["component_id"] == "PAGE-1")
    assert row["category"] == "counted_elsewhere" and row["estimate_status"] == "他へ計上"


def test_summary_mapping_line_shows_split_and_merged():
    """summary.md に原項目→ケース件数と、分解・統合を示す。"""
    from estimator.export import _mapping_lines
    norm = {"source_items": [{"item_id": "A"}, {"item_id": "B"}, {"item_id": "C"}],
            "cases": [{"case_id": "C-1", "source_item_ids": ["A"]}, {"case_id": "C-2", "source_item_ids": ["A"]},
                      {"case_id": "C-3", "source_item_ids": ["B", "C"]}]}
    (line,) = _mapping_lines(norm)
    assert "原項目 3件 → 見積もりケース 3件" in line and "A" in line and "C-3" in line
    assert _mapping_lines(None) == []
