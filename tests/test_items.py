"""大規模資料向けの item_groups / source_item_query、.xlsm、見積もり状態規則、共通Impossible、条件一覧のテスト。"""

import json
import zipfile
from pathlib import Path

from openpyxl import Workbook

from conftest import ADAPTER
from estimator.items import expand_normalized
from estimator.workbook_reader import read_workbook
from helpers import build_cases_standard, build_common_standard, write_config, write_json
from test_pipeline import assessments, design, questions, run_cli

AX = ["AX-01", "AX-02", "AX-03", "AX-04", "AX-05"]


def build_large(path: Path, n_items=12, rows_per_item=3, with_id=True) -> Path:
    """観点×ブラウザーの繰返しを持つ、複数行＋結合セルの人工項目書。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "機能X"
    ws["A1"] = "機能X 項目書"
    ws.append([])
    ws.append(["番号", "観点", "枝番", "操作", "確認", "ブラウザー"])  # R3
    r = 4
    for i in range(n_items):
        view = f"観点{i // 3 + 1}"
        browser = ["Chromium", "Firefox", "WebKit"][i % 3]
        for k in range(rows_per_item):
            ws.cell(r + k, 3, k + 1)
            ws.cell(r + k, 4, f"操作{k + 1}")
            ws.cell(r + k, 5, f"確認{k + 1}")
        if with_id:
            ws.cell(r, 1, f"X-{i + 1:03d}")
        ws.cell(r, 2, view)
        ws.cell(r, 6, browser)
        if rows_per_item > 1:
            for col in (1, 2, 6):
                ws.merge_cells(start_row=r, start_column=col, end_row=r + rows_per_item - 1, end_column=col)
        r += rows_per_item
    wb.save(path)
    return path


def raw_of(path, wid="W1"):
    return {"workbooks": [read_workbook(path, wid)]}


def group(**kw):
    g = {"group_id": "G1", "region_ids": ["W1-S1-R002"], "workbook_id": "W1", "sheet": "機能X", "first_row": 4,
         "last_row": 39, "rows_per_item": 3, "last_column": "F", "id_column": "A", "title_column": "B",
         "disposition": "mapped", "reason": ""}
    g.update(kw)
    return g


def norm_with(groups, cases=()):
    return {"schema_version": 1, "regions": [], "source_items": [], "item_groups": groups, "cases": list(cases)}


def test_group_expansion_multirow(tmp_path):
    raw = raw_of(build_large(tmp_path / "x.xlsx"))
    out, errs = expand_normalized(norm_with([group()]), raw)
    assert errs == []
    items = out["source_items"]
    assert len(items) == 12
    assert items[0]["item_id"] == "X-001" and items[0]["refs"][0]["range"] == "A4:F6"
    assert items[-1]["refs"][0]["range"] == "A37:F39" and items[0]["title"] == "観点1"


def test_group_without_ids_assigns_internal_ids(tmp_path):
    raw = raw_of(build_large(tmp_path / "x.xlsx", with_id=False))
    out, errs = expand_normalized(norm_with([group(id_column=None, item_id_prefix="AUTO-X-")]), raw)
    assert errs == [] and out["source_items"][2]["item_id"] == "AUTO-X-0003"
    assert all(i["id_assigned"] and i["original_id"] is None for i in out["source_items"])
    _, errs = expand_normalized(norm_with([group()]), raw)
    assert errs and errs[0]["code"] == "ITEM_ID_EMPTY"


def test_group_rows_must_divide(tmp_path):
    raw = raw_of(build_large(tmp_path / "x.xlsx"))
    _, errs = expand_normalized(norm_with([group(last_row=38)]), raw)
    assert errs[0]["code"] == "BAD_GROUP"
    _, errs = expand_normalized(norm_with([group(last_row=500)]), raw)
    assert errs[0]["code"] == "BAD_REF"


def test_query_selects_items_across_merged_rows(tmp_path):
    raw = raw_of(build_large(tmp_path / "x.xlsx"))
    case = {"case_id": "C1", "title": "観点1", "source_item_ids": [], "relation": "merged", "summary": "",
            "steps": [], "expected": [], "step_count": 3, "assertion_count": 3, "data_pattern_count": 3,
            "source_item_query": {"group_ids": ["G1"], "where": [{"column": "B", "equals": "観点1"}]}}
    out, errs = expand_normalized(norm_with([group()], [case]), raw)
    assert errs == [] and out["cases"][0]["source_item_ids"] == ["X-001", "X-002", "X-003"]
    case2 = dict(case, case_id="C2", source_item_query={"group_ids": ["G1"], "where": [
        {"column": "F", "equals": "WebKit"}, {"column": "D", "contains": "操作3"}]})
    out, _ = expand_normalized(norm_with([group()], [case2]), raw)
    assert out["cases"][0]["source_item_ids"] == ["X-003", "X-006", "X-009", "X-012"]
    case3 = dict(case, case_id="C3", source_item_query={"group_ids": ["G1"], "where": [{"column": "B", "equals": "無い"}]})
    _, errs = expand_normalized(norm_with([group()], [case3]), raw)
    assert errs[0]["code"] == "EMPTY_QUERY"


def test_xlsm_is_read_and_vba_reported(tmp_path):
    p = build_large(tmp_path / "big.xlsm")
    wb = read_workbook(p)
    assert wb["macro_enabled_format"] and not wb["vba_present"]
    vba = tmp_path / "vba.xlsm"
    with zipfile.ZipFile(p) as zin, zipfile.ZipFile(vba, "w") as zout:
        for n in zin.namelist():
            zout.writestr(n, zin.read(n))
        zout.writestr("xl/vbaProject.bin", b"\0" * 10)
    wb = read_workbook(vba)
    assert wb["vba_present"] and any(w["code"] == "VBA_PRESENT" for w in wb["warnings"])


def _setup_large_run(tmp_path, capsys, run_id="L1"):
    c = build_cases_standard(tmp_path / "cases_v0.1.xlsx")
    m = build_common_standard(tmp_path / "common_v0.1.xlsx")
    big = build_large(tmp_path / "big.xlsm")
    cfg = write_config(tmp_path / "cfg.yaml", c, m, ADAPTER)
    code, out = run_cli(capsys, "init-run", "--input", str(big), "--config", str(cfg), "--run-id", run_id)
    assert code == 0, out
    run_cli(capsys, "extract", "--run", run_id)
    return Path(out["run_dir"])


def test_pipeline_with_groups(tmp_path, capsys):
    rd = _setup_large_run(tmp_path, capsys)
    inv = json.loads((rd / "source_inventory.json").read_text())
    regions = [{"region_id": r["region_id"], "role": "test_table" if r["row_count"] > 2 else "heading",
                "disposition": "parsed" if r["row_count"] > 2 else "not_test", "reason": "r"} for r in inv["regions"]]
    table = next(r for r in inv["regions"] if r["row_count"] > 2)
    cases = [{"case_id": f"C{v}", "title": f"観点{v}", "source_item_ids": [], "relation": "merged",
              "relation_reason": "ブラウザー違いを1つのパラメータ化ケースへ", "summary": "", "steps": [], "expected": [],
              "step_count": 3, "assertion_count": 3, "data_pattern_count": 3,
              "source_item_query": {"group_ids": ["G1"], "where": [{"column": "B", "equals": f"観点{v}"}]}}
             for v in range(1, 5)]
    write_json(rd / "ai/normalized_cases.json", {"schema_version": 1, "regions": regions, "source_items": [],
                                                 "item_groups": [group(region_ids=[table["region_id"]])],
                                                 "cases": cases})
    code, out = run_cli(capsys, "validate-stage", "--run", "L1", "--stage", "normalize")
    assert code == 0, out
    d = design()
    for comp in d["components"]:
        comp["used_by"] = ["C1", "C2", "C3", "C4"]
    write_json(rd / "ai/common_design.json", d)
    a = assessments()
    a["case_assessments"] = [dict(a["case_assessments"][0], case_id=f"C{v}") for v in range(1, 5)]
    write_json(rd / "ai/assessments.json", a)
    write_json(rd / "ai/questions.json", questions())
    for st in ("common_design", "assessments"):
        code, out = run_cli(capsys, "validate-stage", "--run", "L1", "--stage", st)
        assert code == 0, out
    assert run_cli(capsys, "calculate", "--run", "L1")[0] == 0
    code, out = run_cli(capsys, "validate", "--run", "L1")
    assert code == 0, out
    assert run_cli(capsys, "export", "--run", "L1")[0] == 0
    code, out = run_cli(capsys, "profile", "--run", "L1", "--sheet", "機能X", "--rows", "4-39", "--columns", "B,F",
                        "--group-by", "B")
    assert out["columns"]["F"]["distinct"] == 3 and out["group_by"]["distinct"] == 4


# ---------------------------------------------------------------- 見積もり状態規則・共通Impossible・条件一覧

def _axes(*s, assumptions=None):
    d = {a: {"score": v, "condition": "c", "evidence": "e"} for a, v in zip(AX, s)}
    if assumptions:
        d["AX-05"]["assumptions"] = assumptions
    return d


def _run_basic(tmp_path, capsys):
    from test_pipeline import normalized
    from helpers import build_customer

    c = build_cases_standard(tmp_path / "cases_v0.1.xlsx")
    m = build_common_standard(tmp_path / "common_v0.1.xlsx")
    cust = build_customer(tmp_path / "cust.xlsx")
    cfg = write_config(tmp_path / "cfg.yaml", c, m, ADAPTER)
    code, out = run_cli(capsys, "init-run", "--input", str(cust), "--config", str(cfg), "--run-id", "S1")
    run_cli(capsys, "extract", "--run", "S1")
    rd = Path(out["run_dir"])
    write_json(rd / "ai/normalized_cases.json", normalized())
    write_json(rd / "ai/common_design.json", design())
    write_json(rd / "ai/questions.json", questions())
    return rd


def _codes(capsys, stage="assessments"):
    code, out = run_cli(capsys, "validate-stage", "--run", "S1", "--stage", stage)
    return {e["code"] for e in out["errors"]}


def test_status_rule_rule07(tmp_path, capsys):
    rd = _run_basic(tmp_path, capsys)
    a = assessments()
    a["case_assessments"][0]["axes"] = _axes(1, 1, 1, 2, 1)  # AX-04=2 → 条件付きのはず
    write_json(rd / "ai/assessments.json", a)
    assert "STATUS_RULE" in _codes(capsys)
    a["case_assessments"][0].update(estimate_status="条件付き", conditions=["期待値の文言を仮定"])
    write_json(rd / "ai/assessments.json", a)
    assert "STATUS_RULE" not in _codes(capsys)
    a["case_assessments"][0].update(axes=_axes(1, 1, 1, 1, 1, assumptions=["固定データあり"]),
                                    estimate_status="確定可能", conditions=[])
    write_json(rd / "ai/assessments.json", a)
    assert "STATUS_RULE" in _codes(capsys)  # 仮定あり → 条件付き
    a["case_assessments"][0].update(axes=_axes(1, 1, 1, 1, 1), feasibility="impossible_confirmed",
                                    impossible_reason="x", estimate_status="要確認")
    write_json(rd / "ai/assessments.json", a)
    assert "STATUS_RULE" in _codes(capsys)  # RULE-03


def test_common_impossible(tmp_path, capsys):
    rd = _run_basic(tmp_path, capsys)
    a = assessments()
    a["common_assessments"][1].update(difficulty="Impossible")
    write_json(rd / "ai/assessments.json", a)
    assert "NO_REASON" in _codes(capsys)
    a["common_assessments"][1].update(impossible_reason="合意範囲で手段なしを確認済")
    write_json(rd / "ai/assessments.json", a)
    assert _codes(capsys) == set()
    code, out = run_cli(capsys, "calculate", "--run", "S1")
    arch = out["architectures"][0]
    assert arch["counts"]["components"]["impossible"] == 1
    assert arch["subtotals"]["common"] == {"low": 1.0, "std": 2.0, "high": 3.0}  # BASEのみ
    assert arch["estimate_completeness"] == "一部未算定"


def test_conditions_listed_in_summary(tmp_path, capsys):
    rd = _run_basic(tmp_path, capsys)
    a = assessments()
    a["case_assessments"][0].update(axes=_axes(2, 1, 1, 2, 1), estimate_status="条件付き",
                                    conditions=["一覧の表示件数を仮定"])
    a["common_assessments"][0].update(estimate_status="条件付き", conditions=["CIなしを仮定"])
    write_json(rd / "ai/assessments.json", a)
    for st in ("normalize", "common_design", "assessments"):
        assert run_cli(capsys, "validate-stage", "--run", "S1", "--stage", st)[0] == 0
    run_cli(capsys, "calculate", "--run", "S1")
    run_cli(capsys, "validate", "--run", "S1")
    assert run_cli(capsys, "export", "--run", "S1")[0] == 0
    md = (rd / "summary.md").read_text()
    assert "条件付きの仮定・条件 (2件)" in md and "一覧の表示件数を仮定" in md and "CIなしを仮定" in md
    from openpyxl import load_workbook
    ws = load_workbook(rd / "estimate.xlsx")["サマリー"]
    assert any(c.value == "一覧の表示件数を仮定" for row in ws.iter_rows() for c in row)
