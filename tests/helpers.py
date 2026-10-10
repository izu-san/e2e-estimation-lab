"""テスト用の人工データ生成。

ここで作る基準Excelは「アダプターの読取りと計算を検証するための人工物」であり、
提供される正本の基準Excelではない。値は計算テスト用の任意値。
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook

COMMON_TYPES = ["BASE", "AUTH", "FIX", "DATA", "PAGE", "COMP", "FLOW", "UTIL", "EXT", "RUN", "CI", "REPORT", "DOC", "POC"]


def build_cases_standard(path: Path, *, low=(1, 2, 3), mid=(3, 5, 8), high=(8, 16, 24),
                         low_max=5, high_min=8, version="v0.1", extra_text=None, dup_mid=False,
                         formula_threshold=False) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "表紙"
    ws["A1"] = "Playwright 自動化工数標準設定（テスト用人工データ）"
    ws["A3"], ws["B3"] = "版", version
    if extra_text:
        ws["A5"] = extra_text
    ax = wb.create_sheet("判定基準")
    ax.append(["軸ID", "判定軸", "1点", "2点", "3点"])
    for i, name in enumerate(["操作手順", "確認手順", "実装手段", "あいまい度", "データ準備・後処理"], 1):
        ax.append([f"AX-0{i}", name, "単純", "中程度", "複雑"])
    ef = wb.create_sheet("工数標準")
    ef.append(["総合難易度の閾値"])
    ef.append(["Low上限", low_max, "点"])
    if formula_threshold:
        ef.append(["High閾値", "=B2+3", "点"])  # 保存済み計算結果なし
    else:
        ef.append(["High閾値", high_min, "点"])
    ef.append([])
    ef.append(["難易度", "下限", "標準", "上限", "説明"])
    ef.append(["Low", *low, "単純"])
    ef.append(["Mid", *mid, "中程度"])
    if dup_mid:
        ef.append(["Mid", *mid, "重複"])
    ef.append(["High", *high, "複雑"])
    ef.append(["Impossible", None, None, None, "工数なし"])
    wb.save(path)
    return path


def default_common_rows():
    rows = []
    for i, t in enumerate(COMMON_TYPES):
        for d, mult in (("Low", 1), ("Mid", 2), ("High", 4)):
            rows.append([t, d, "1式" if t in ("BASE", "CI", "RUN") else "1部品", mult * 1.0, mult * 2.0, mult * 3.0])
    return rows


def build_common_standard(path: Path, *, rows=None, version="v0.1", example=None) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "概要"
    ws["A1"] = "Playwright 汎用共通部分工数基準（テスト用人工データ）"
    ws["A2"], ws["B2"] = "版", version
    ef = wb.create_sheet("工数表")
    ef.append(["種別ID", "難易度", "単位", "下限", "標準", "上限", "判定条件"])
    for r in rows if rows is not None else default_common_rows():
        ef.append([*r, "条件文"])
    if example is not None:
        ex = wb.create_sheet("判定集計例")
        ex.append(["例の説明"])
        ex.append(["種別ID", "難易度", "数量", "状態"])
        for r in example:
            ex.append(r)
    wb.save(path)
    return path


def write_config(path: Path, cases: Path, common: Path, adapter: Path, **extra) -> Path:
    import yaml

    cfg = {"schema_version": 1,
           "standards": {"cases": str(cases), "common": str(common), "adapter": str(adapter)},
           "execution": {"site_interaction": False, "max_repair_attempts": 2}}
    cfg.update(extra)
    path.write_text(yaml.safe_dump(cfg, allow_unicode=True), encoding="utf-8")
    return path


def build_customer(path: Path, *, col_order=("ID", "前提", "手順", "期待結果"), sheet="テスト項目",
                   with_desc=True, header_names=None) -> Path:
    """小さな人工の顧客項目書。列順・列名・説明行・結合セル・非表示行を含む。"""
    names = header_names or {}
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    r = 1
    if with_desc:
        ws.cell(1, 1, "本書は人工のテスト項目書です。")
        ws.cell(2, 1, "=SUM(1,2)")  # 保存値なしの数式
        r = 4
    headers = [names.get(c, c) for c in col_order]
    for j, h in enumerate(headers, 1):
        ws.cell(r, j, h)
    data = [
        {"ID": "T-1", "前提": "ログイン済み", "手順": "1. 商品一覧を開く", "期待結果": "一覧が表示される"},
        {"ID": None, "前提": None, "手順": "2. 商品を選ぶ", "期待結果": "詳細が表示される"},
        {"ID": "T-2", "前提": "未ログイン", "手順": "=ログインボタンを押す", "期待結果": "ログイン画面"},
    ]
    for k, d in enumerate(data, 1):
        for j, c in enumerate(col_order, 1):
            v = d[c]
            if v is not None:
                cell = ws.cell(r + k, j, v)
                if isinstance(v, str) and v.startswith("="):
                    cell.data_type = "s"  # 数式ではなく「=」で始まる文字列
    id_col = col_order.index("ID") + 1
    pre_col = col_order.index("前提") + 1
    ws.merge_cells(start_row=r + 1, start_column=id_col, end_row=r + 2, end_column=id_col)
    ws.merge_cells(start_row=r + 1, start_column=pre_col, end_row=r + 2, end_column=pre_col)
    ws.row_dimensions[r + 3].hidden = True
    ws.cell(r + 6, 1, "参考: 対象URLは別紙")
    hidden = wb.create_sheet("非表示メモ")
    hidden["A1"] = "旧版の項目"
    hidden.sheet_state = "hidden"
    wb.save(path)
    return path


def write_json(path: Path, obj) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")
