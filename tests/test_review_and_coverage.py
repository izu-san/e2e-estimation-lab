"""runs/ 横断調査（2026-10）で見つかった再発パターンへの対策のテスト。

- 意味レビューの指摘に対応の記録を求め、未対応を状態と summary に出す
- 計算規則の版が変わったことを再計算照合で区別する
- 未確定部分を疑似ケースにせず pending_parts に保留する
- 共通部品が所有する作業をケースの軸で再加点していないか（警告）
- 資料情報・前提の行が質問・仮定から漏れていないか（警告）
"""

import json

from helpers import write_json
from test_pipeline import _stage, assessments, design, normalized, questions, run, run_cli  # noqa: F401

FINDING = {"target": "C-001", "related": ["C-002"], "kind": "scoring", "message": "根拠が弱い", "suggestion": "再確認"}


def _through_export(capsys):
    for st in ("normalize", "screens", "common_design", "questions", "assessments"):
        assert run_cli(capsys, "validate-stage", "--run", "t1", "--stage", st)[0] == 0
    assert run_cli(capsys, "calculate", "--run", "t1")[0] == 0
    assert run_cli(capsys, "validate", "--run", "t1")[0] == 0
    assert run_cli(capsys, "export", "--run", "t1")[0] == 0


def _review(rd, *findings):
    write_json(rd / "ai/semantic_review.json", {"schema_version": 1, "reviewer": "ai", "findings": list(findings)})


def test_review_findings_need_resolution_before_mark_review(run, capsys):
    rd, _ = run
    _through_export(capsys)
    _review(rd, FINDING)
    code, errors, warns = _stage(capsys, "semantic_review")
    assert code == 0 and "FINDING_UNRESOLVED" in warns
    code, out = run_cli(capsys, "mark-review", "--run", "t1", "--by", "ai", "--context", "independent")
    assert code != 0 and "resolution" in out["message"]
    # 指摘ファイルがあっても記録前は「済」にしない
    assert json.loads((rd / "validation.json").read_text())["status"]["semantic_review_status"] == "未実施"
    _review(rd, {**FINDING, "resolution": {"status": "question", "note": "質問にした", "question_ids": ["Q-404"]}})
    code, errors, _ = _stage(capsys, "semantic_review")
    assert code == 1 and "UNKNOWN_QUESTION" in errors
    _review(rd, {**FINDING, "resolution": {"status": "question", "note": "質問にした"}})
    assert "NO_QUESTION" in _stage(capsys, "semantic_review")[1]


def test_deferred_findings_show_in_status_and_summary(run, capsys):
    rd, _ = run
    _through_export(capsys)
    _review(rd, {**FINDING, "resolution": {"status": "deferred", "note": "基準担当の回答待ち"}},
            {**FINDING, "target": "C-002", "resolution": {"status": "fixed", "note": "AX-01を1点に修正し再計算"}},
            {**FINDING, "target": "overall", "resolution": {"status": "question", "note": "起票", "question_ids": ["Q1"]}})
    code, out = run_cli(capsys, "mark-review", "--run", "t1", "--by", "ai", "--context", "self")
    assert code == 0, out
    assert out["semantic_review_status"] == "AIレビュー済（自己レビュー・未対応1件）"
    summary = (rd / "summary.md").read_text(encoding="utf-8")
    assert "## 意味レビューの指摘 (3件、うち未対応1件)" in summary
    assert summary.index("[未対応]") < summary.index("[修正済]")  # 未対応を先に出す
    from openpyxl import load_workbook
    ws = load_workbook(rd / "estimate.xlsx")["意味レビュー"]
    assert [c.value for c in ws[1]][:6] == ["対象", "関連", "種類", "指摘", "修正案", "対応"]
    # レビュー記録後に採点が変わったら、状態に出す
    a = assessments()
    a["case_assessments"][0]["axes"]["AX-02"]["score"] = 2
    write_json(rd / "ai/assessments.json", a)
    assert run_cli(capsys, "validate-stage", "--run", "t1", "--stage", "assessments")[0] == 0
    assert run_cli(capsys, "calculate", "--run", "t1")[0] == 0
    code, out = run_cli(capsys, "validate", "--run", "t1")
    assert "レビュー記録後に見積もり変更" in out["status"]["semantic_review_status"]


def test_rule_version_change_is_reported_separately(run, capsys):
    rd, _ = run
    _through_export(capsys)
    est = json.loads((rd / "estimate.json").read_text())
    est["rule_versions"] = {"cases": "case-rule-0", "common": "common-rule-0"}
    est["architectures"][0]["components"][0]["estimate_status"] = "要確認"
    (rd / "estimate.json").write_text(json.dumps(est, ensure_ascii=False))
    code, out = run_cli(capsys, "validate", "--run", "t1")
    assert code == 1 and [e["code"] for e in out["errors"]] == ["RULE_VERSION_CHANGED"]
    assert out["status"]["execution_status"] == "処理失敗"


def test_pending_parts_count_as_unresolved(run, capsys):
    rd, _ = run
    n = normalized()
    n["source_items"][0]["pending_parts"] = [{"summary": "スマホ表示の確認（仕様未提示）", "question_ids": ["Q1"],
                                              "refs": [{"workbook_id": "W1", "sheet": "テスト項目", "range": "D6"}]}]
    write_json(rd / "ai/normalized_cases.json", n)
    _through_export(capsys)
    est = json.loads((rd / "estimate.json").read_text())
    a = est["architectures"][0]
    assert a["counts"]["pending_parts"] == 1
    assert a["counts"]["unresolved_items"] == 2  # C-002（要確認）+ 保留部分
    assert "ケースにせず保留した未確定部分 1件" in (rd / "summary.md").read_text(encoding="utf-8")
    n["source_items"][0]["pending_parts"][0]["question_ids"] = ["Q-404"]
    write_json(rd / "ai/normalized_cases.json", n)
    assert "UNKNOWN_QUESTION" in _stage(capsys, "questions")[1]
    n["source_items"][0]["pending_parts"][0]["question_ids"] = []
    write_json(rd / "ai/normalized_cases.json", n)
    assert "SCHEMA" in _stage(capsys, "normalize")[1]  # 保留には質問が必須


def test_pseudo_case_split_off_as_pending_warns(run, capsys):
    rd, _ = run
    n = normalized()
    c1 = n["cases"][0]
    c1["relation"] = "split"
    c1["relation_reason"] = "確定部分と未提示部分"
    c1["step_count"] = 1
    n["cases"].append({**c1, "case_id": "C-001B", "title": "未提示部分", "steps": ["確認"], "step_count": 1})
    write_json(rd / "ai/normalized_cases.json", n)
    d = design()
    d["components"][0]["used_by"].append("C-001B")
    write_json(rd / "ai/common_design.json", d)
    a = assessments()
    a["case_assessments"].append({**a["case_assessments"][1], "case_id": "C-001B", "question_ids": ["Q1"]})
    write_json(rd / "ai/assessments.json", a)
    q = questions()
    q["items"][0]["targets"].append("C-001B")
    write_json(rd / "ai/questions.json", q)
    for st in ("normalize", "screens", "common_design", "questions"):
        assert _stage(capsys, st)[0] == 0
    code, errors, warns = _stage(capsys, "assessments")
    assert code == 0 and "PENDING_SPLIT_CASE" in warns


def test_axis_overlapping_component_needs_note(run, capsys):
    rd, _ = run
    d = design()
    d["components"].append({**d["components"][0], "component_id": "DATA-1", "type_id": "DATA", "name": "ユーザー作成",
                            "elements": ["USER"], "cost_owner": "DATA-1", "used_by": ["C-001"]})
    write_json(rd / "ai/common_design.json", d)
    a = assessments()
    c1 = a["case_assessments"][0]
    c1["uses_components"].append("DATA-1")
    c1["axes"]["AX-05"]["score"] = 2
    a["common_assessments"].append({**a["common_assessments"][0], "component_id": "DATA-1"})
    write_json(rd / "ai/assessments.json", a)
    code, errors, warns = _stage(capsys, "assessments")
    assert code == 0 and "AXIS_COMPONENT_OVERLAP" in warns
    c1["axes"]["AX-05"]["overlap_note"] = "ケース固有の購読解除データの指定"
    write_json(rd / "ai/assessments.json", a)
    assert "AXIS_COMPONENT_OVERLAP" not in _stage(capsys, "assessments")[2]
    c1["axes"]["AX-05"]["score"] = 1  # 部品を呼ぶだけの点数なら注記不要
    del c1["axes"]["AX-05"]["overlap_note"]
    write_json(rd / "ai/assessments.json", a)
    assert "AXIS_COMPONENT_OVERLAP" not in _stage(capsys, "assessments")[2]


def test_premise_rows_must_be_traced(run, capsys):
    rd, _ = run
    inv = json.loads((rd / "source_inventory.json").read_text())
    regs = {r["region_id"]: r for r in inv["regions"]}
    code, errors, warns = _stage(capsys, "questions")
    assert code == 0 and "PREMISE_ROW_UNTRACED" in warns
    q = questions()
    r1, r3 = regs["W1-S1-R001"], regs["W1-S1-R003"]
    q["items"][0]["refs"] = [{"workbook_id": "W1", "sheet": r1["sheet"], "range": r1["range"]}]
    write_json(rd / "ai/questions.json", q)
    assert "PREMISE_ROW_UNTRACED" in _stage(capsys, "questions")[2]  # 参考行（R003）が残る
    q["premise_reviews"] = [{"ref": {"workbook_id": "W1", "sheet": r3["sheet"], "range": r3["range"]},
                             "reason": "別紙URLの案内で見積もりに影響しない"}]
    write_json(rd / "ai/questions.json", q)
    assert "PREMISE_ROW_UNTRACED" not in _stage(capsys, "questions")[2]


def test_consistency_across_formats_and_reused_runs(run, capsys, tmp_path):
    """同じ内容の項目書どうしの結果のばらつき・原IDごとの食い違い・AI成果物を流用したrunを検出する。"""
    import shutil

    from estimator.evaluate import evaluate_consistency
    from helpers import build_customer

    rd1, (_, _, cust) = run
    _through_export(capsys)
    cfg = tmp_path / "cfg.yaml"
    cust2 = build_customer(tmp_path / "cust2.xlsx", header_names={"手順": "操作"})

    def new_run(rid, inp, assess):
        code, out = run_cli(capsys, "init-run", "--input", str(inp), "--config", str(cfg), "--run-id", rid)
        assert code == 0, out
        assert run_cli(capsys, "extract", "--run", rid)[0] == 0
        rd = rd1.parent / rid
        for f in ("normalized_cases", "screens", "common_design", "questions"):
            shutil.copy(rd1 / f"ai/{f}.json", rd / f"ai/{f}.json")
        write_json(rd / "ai/assessments.json", assess)
        assert run_cli(capsys, "calculate", "--run", rid)[0] == 0
        return rd

    a = assessments()
    a["case_assessments"][0]["axes"]["AX-01"]["score"] = 3  # 書式違いで C-001（T-1）が High になった想定
    new_run("t2", cust2, a)
    new_run("t3", cust, assessments())  # t1 のAI成果物を流用した再実行

    res = evaluate_consistency({"name": "g", "inputs": [str(cust), str(cust2)],
                                "tolerances": {"max_relative_spread": {"computed_total_std": 0.0}}})
    assert res["status"] == "failed"
    assert res["spread"]["cases_std"]["values"] == {"cust": 5.0, "cust2": 16.0}
    assert res["item_difficulty_disagreements"] == {"T-1": {"cust": ["Mid"], "cust2": ["High"]}}
    first = next(i for i in res["inputs"] if i["label"] == "cust")
    assert [t["reused_ai_outputs_from"] for t in first["runs"]] == [None, "t1"]
    assert first["independent_runs"] == 1
    assert evaluate_consistency({"name": "g", "inputs": [str(cust), str(cust2)]})["status"] == "report_only"


def test_ax01_screen_rule(run, capsys):
    """操作した画面の数・要因から AX-01 の期待点を求めて照合する（個別基準 COST-13 の仮ルール）。"""
    rd, _ = run
    a = assessments()
    ax = a["case_assessments"][0]["axes"]["AX-01"]  # C-001: AX-01=2
    assert "AX01_NO_SCREENS" in _stage(capsys, "assessments")[2]
    ax["operation_screens"] = ["PG-List"]
    write_json(rd / "ai/assessments.json", a)
    assert "AX01_SCREEN_RULE" in _stage(capsys, "assessments")[2]  # 1画面なら1点
    ax["factors"] = ["dialog"]
    write_json(rd / "ai/assessments.json", a)
    assert "AX01_SCREEN_RULE" not in _stage(capsys, "assessments")[2]  # ダイアログがあれば2点
    del ax["factors"]
    ax["operation_screens"] = ["PG-List", "PG-Cart"]
    write_json(rd / "ai/assessments.json", a)
    _, errors, warns = _stage(capsys, "assessments")
    assert "AX01_SCREEN_RULE" not in warns and "SCREEN_NOT_IN_INVENTORY" in errors  # PG-Cart は画面一覧に無い
    ax["operation_screens"] = ["PG-List", "PG-Detail"]
    write_json(rd / "ai/assessments.json", a)
    _, errors, warns = _stage(capsys, "assessments")
    assert "SCREEN_NOT_IN_INVENTORY" not in errors and "SCREEN_NOT_OPERATED" in warns  # 一覧では PG-Detail は C-002 が操作


def test_screen_inventory(run, capsys):
    """画面一覧を先に確定し、PAGEの構成要素をそれと一致させる（共通基準 BOUND-15）。"""
    from test_pipeline import screens

    rd, _ = run
    assert _stage(capsys, "screens")[0] == 0
    assert not _stage(capsys, "common_design")[1]
    s = screens()
    s["screens"].append({**s["screens"][0], "screen_id": "PG-Cart", "operated_by": ["C-009"], "checked_by": []})
    s["screens"][0]["checked_by"] = ["C-001"]
    s["screens"][1]["basis"] = "observed"
    s["not_counted"] = [{"name": "ブランド別一覧", "reason": "同じURLのフィルター", "merged_into": "PG-Brand"}]
    write_json(rd / "ai/screens.json", s)
    _, errors, _ = _stage(capsys, "screens")
    assert {"UNKNOWN_CASE", "SCREEN_ROLE_OVERLAP", "OBSERVED_WITHOUT_SURVEY", "UNKNOWN_SCREEN"} <= errors
    s = screens()
    s["screens"][1]["operated_by"] = []
    write_json(rd / "ai/screens.json", s)
    assert "CASE_NO_SCREEN" in _stage(capsys, "screens")[2]
    s = screens()
    s["screens"].append({**s["screens"][0], "screen_id": "PG-Cart"})
    write_json(rd / "ai/screens.json", s)
    assert "SCREEN_NOT_IN_PAGE" in _stage(capsys, "common_design")[1]  # PAGEに PG-Cart が無い
    write_json(rd / "ai/screens.json", screens())
    d = design()
    d["components"][1]["elements"] = ["PG-List", "PG-Detail", "PG-Extra"]
    write_json(rd / "ai/common_design.json", d)
    assert "PAGE_NOT_IN_SCREENS" in _stage(capsys, "common_design")[1]
    (rd / "ai/screens.json").unlink()
    assert "SCREENS_MISSING" in _stage(capsys, "common_design")[1]


def test_component_rules(run, capsys):
    """必須の種別（RUN）と、FLOW・COMP の最小利用数（共通基準 BOUND-16〜18 の仮ルール）。"""
    rd, _ = run
    d = design()
    base = d["components"][0]
    d["components"].append({**base, "component_id": "FLOW-1", "type_id": "FLOW", "name": "業務フロー",
                            "elements": ["FLOW"], "cost_owner": "FLOW-1", "used_by": ["C-001"], "depends_on": ["PAGE-1"]})
    write_json(rd / "ai/common_design.json", d)
    code, errors, warns = _stage(capsys, "common_design")
    assert code == 0 and {"REQUIRED_COMPONENT_MISSING", "COMPONENT_SINGLE_USE"} <= warns
    d["components"][-1]["used_by"] = ["C-001", "C-002"]
    d["components"].append({**base, "component_id": "RUN-1", "type_id": "RUN", "name": "全体実行",
                            "elements": ["RUN"], "cost_owner": "RUN-1"})
    write_json(rd / "ai/common_design.json", d)
    warns = _stage(capsys, "common_design")[2]
    assert not {"REQUIRED_COMPONENT_MISSING", "COMPONENT_SINGLE_USE"} & warns


def test_page_difficulty_rule(run, capsys):
    """PAGEの難易度を画面の特徴から照合する（共通基準 BOUND-19 の仮ルール）。"""
    rd, _ = run
    a = assessments()
    page = a["common_assessments"][1]  # PAGE-1: Mid, elements PG-List / PG-Detail
    prof = lambda el, **kw: {"element": el, "inputs": 1, "dynamic_list": False, "dialog": False, "file": False,  # noqa: E731
                             "complex_ui": False, **kw}
    assert "PAGE_NO_PROFILE" in _stage(capsys, "assessments")[2]
    page["screen_profiles"] = [prof("PG-List", dynamic_list=True), prof("PG-Detail")]
    write_json(rd / "ai/assessments.json", a)
    assert "PAGE_MIXED_DIFFICULTY" in _stage(capsys, "assessments")[2]  # Mid と Low の混在
    page["screen_profiles"] = [prof("PG-List", dynamic_list=True), prof("PG-Detail", inputs=6)]
    write_json(rd / "ai/assessments.json", a)
    warns = _stage(capsys, "assessments")[2]
    assert not {"PAGE_MIXED_DIFFICULTY", "PAGE_DIFFICULTY_RULE", "PAGE_PROFILE_MISMATCH"} & warns
    page["difficulty"] = "Low"
    write_json(rd / "ai/assessments.json", a)
    assert "PAGE_DIFFICULTY_RULE" in _stage(capsys, "assessments")[2]
    page["screen_profiles"] = [prof("PG-List")]
    write_json(rd / "ai/assessments.json", a)
    assert "PAGE_PROFILE_MISMATCH" in _stage(capsys, "assessments")[2]
