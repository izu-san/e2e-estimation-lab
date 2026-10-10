"""構造抽出テスト（10.2の機械的部分）。意味判断はAIの担当なので、ここでは原文・構造の保持を確認する。"""

import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from estimator.export import _cell
from estimator.workbook_reader import UnsupportedWorkbook, detect_regions, read_workbook
from helpers import build_customer


def _sheet(wb, name):
    return next(s for s in wb["sheets"] if s["name"] == name)


def test_merged_cells_not_forward_filled(tmp_path):
    wb = read_workbook(build_customer(tmp_path / "a.xlsx"))
    sh = _sheet(wb, "テスト項目")
    refs = {c["ref"]: c for c in sh["cells"]}
    assert refs["A5"]["value"] == "T-1"
    assert "A6" not in refs  # 結合範囲の非代表セルは空のまま
    assert sh["merged_cover"]["A6"] == "A5:A6"
    assert {"range": "A5:A6", "anchor": "A5"} in sh["merged"]


def test_hidden_and_formula_preserved(tmp_path):
    wb = read_workbook(build_customer(tmp_path / "a.xlsx"))
    sh = _sheet(wb, "テスト項目")
    assert sh["hidden_rows"] == [7]
    assert _sheet(wb, "非表示メモ")["state"] == "hidden"
    a2 = next(c for c in sh["cells"] if c["ref"] == "A2")
    assert a2["type"] == "formula" and a2["cached_missing"] is True and a2["cached_value"] is None
    c7 = next(c for c in sh["cells"] if c["ref"] == "C7")
    assert c7["type"] == "string" and c7["value"].startswith("=")
    codes = {w["code"] for w in wb["warnings"]}
    assert {"FORMULA_NO_CACHE", "HIDDEN_SHEET", "HIDDEN_ROWS_COLS"} <= codes


@pytest.mark.parametrize("kwargs", [
    {},
    {"col_order": ("期待結果", "手順", "ID", "前提")},                         # 列順変更
    {"header_names": {"ID": "No.", "手順": "操作", "期待結果": "確認内容"}},   # 列名変更
    {"sheet": "Sheet_別名", "with_desc": False},                             # シート名変更・説明行なし
])
def test_region_detection_is_template_independent(tmp_path, kwargs):
    wb = read_workbook(build_customer(tmp_path / "renamed_input.xlsx", **kwargs), "WX")
    sh = wb["sheets"][0]
    regs = detect_regions(sh, "WX")
    table = [r for r in regs if r["row_count"] == 4]
    assert len(table) == 1  # 見出し+3行の表として1領域
    assert table[0]["nonempty_cells"] == 14
    assert all(r["region_id"].startswith("WX-S1-") for r in regs)


def test_images_detected(tmp_path):
    from openpyxl.drawing.image import Image
    from PIL import Image as PILImage

    PILImage.new("RGB", (10, 10)).save(tmp_path / "p.png")
    wb = Workbook()
    wb.active["A1"] = "画面イメージ"
    wb.active.add_image(Image(str(tmp_path / "p.png")), "C3")
    wb.save(tmp_path / "img.xlsx")
    r = read_workbook(tmp_path / "img.xlsx")
    assert r["sheets"][0]["drawings"] == [{"kind": "image", "anchor": "C3", "text": None, "part": "xl/drawings/drawing1.xml"}]
    regs = detect_regions(r["sheets"][0], "W1")
    assert any(x.get("kind") == "drawing" for x in regs)


def test_comments_detected(tmp_path):
    from openpyxl.comments import Comment

    wb = Workbook()
    wb.active["B2"] = "x"
    wb.active["B2"].comment = Comment("仕様未確定", "顧客")
    wb.save(tmp_path / "c.xlsx")
    assert read_workbook(tmp_path / "c.xlsx")["sheets"][0]["comments"][0]["text"] == "仕様未確定"


def test_unsupported_formats(tmp_path):
    ole = tmp_path / "old.xls"
    ole.write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + b"\0" * 100)
    with pytest.raises(UnsupportedWorkbook, match="OLE"):
        read_workbook(ole)
    enc = tmp_path / "locked.xlsx"
    enc.write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + b"\0" * 100)
    with pytest.raises(UnsupportedWorkbook, match="暗号化"):
        read_workbook(enc)
    bad = tmp_path / "broken.xlsx"
    bad.write_bytes(b"not a zip")
    with pytest.raises(UnsupportedWorkbook):
        read_workbook(bad)
    csv = tmp_path / "a.csv"
    csv.write_text("a,b")
    with pytest.raises(UnsupportedWorkbook, match="正式対応外"):
        read_workbook(csv)


def test_export_does_not_create_formulas(tmp_path):
    wb = Workbook()
    ws = wb.active
    _cell(ws, 1, 1, "=HYPERLINK(\"http://x\")")
    _cell(ws, 2, 1, None)
    _cell(ws, 3, 1, 0)
    wb.save(tmp_path / "o.xlsx")
    with zipfile.ZipFile(tmp_path / "o.xlsx") as z:
        assert b"<f>" not in z.read("xl/worksheets/sheet1.xml")
    ws2 = load_workbook(tmp_path / "o.xlsx").active
    assert ws2["A1"].data_type == "s" and ws2["A1"].value.startswith("=")
    assert ws2["A2"].value is None and ws2["A3"].value == 0
