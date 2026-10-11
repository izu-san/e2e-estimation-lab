"""estimate.xlsx / summary.md の出力。値はすべて estimate.json（スクリプト算出）から転記する。"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .io_utils import atomic_save_workbook, read_json, write_text
from .items import load_normalized

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
_FORMULA_PREFIX = ("=", "+", "-", "@", "\t", "\r")


def _cell(ws, row: int, col: int, value: Any):
    """原文が = 等で始まってもExcel数式として解釈させない。Noneは空欄（0にしない）。"""
    if isinstance(value, (list, tuple)):
        value = "\n".join(str(v) for v in value)
    elif isinstance(value, dict):
        value = "; ".join(f"{k}={v}" for k, v in value.items())
    c = ws.cell(row=row, column=col)
    if value is None:
        return c
    c.value = value
    if isinstance(value, str):
        c.data_type = "s"
        if value.startswith(_FORMULA_PREFIX):
            c.quotePrefix = True
    return c


def _table(ws, headers: list[str], rows: list[list[Any]], widths: dict[int, int] | None = None):
    for j, h in enumerate(headers, start=1):
        c = _cell(ws, 1, j, h)
        c.font = Font(bold=True)
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, r in enumerate(rows, start=2):
        for j, v in enumerate(r, start=1):
            c = _cell(ws, i, j, v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    for j in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(j)].width = (widths or {}).get(j, 16)


def conditional_items(assess: dict) -> list[dict]:
    """条件付きの個別ケース・共通作業と、その仮定・条件の一覧（サマリー表示用）。"""
    out = []
    for a in assess["case_assessments"]:
        if a["estimate_status"] != "条件付き":
            continue
        assumptions = [f"{ax}: {x}" for ax, v in a["axes"].items() for x in v.get("assumptions", [])]
        out.append({"kind": "個別ケース", "id": a["case_id"], "arch_ids": a["arch_ids"],
                    "conditions": list(a.get("conditions", [])), "assumptions": assumptions})
    for c in assess["common_assessments"]:
        if c["estimate_status"] == "条件付き":
            out.append({"kind": "共通作業", "id": c["component_id"], "arch_ids": [],
                        "conditions": list(c.get("conditions", [])), "assumptions": []})
    return out


def _ref_str(refs: list[dict]) -> str:
    return "\n".join(f"{r['workbook_id']}:{r['sheet']}!{r['range']}" for r in refs)


def export(rd: Path) -> dict:
    est = read_json(rd / "estimate.json")
    val = read_json(rd / "validation.json")
    man = read_json(rd / "manifest.json")
    norm = load_normalized(rd)
    design = read_json(rd / "ai/common_design.json")
    assess = read_json(rd / "ai/assessments.json")
    qpath = rd / "ai/questions.json"
    questions = read_json(qpath)["items"] if qpath.exists() else []
    inv = read_json(rd / "source_inventory.json")
    case_by_id = {c["case_id"]: c for c in norm["cases"]}
    comp_by_id = {c["component_id"]: c for c in design["components"]}
    ca_by_id = {c["component_id"]: c for c in assess["common_assessments"]}
    item_by_id = {i["item_id"]: i for i in norm["source_items"]}
    status = val["status"]

    wb = Workbook()
    ws = wb.active
    ws.title = "サマリー"
    rows: list[list[Any]] = [
        ["run_id", man["run_id"]],
        ["入力", ", ".join(i["file_name"] for i in man["inputs"])],
        ["execution_status", status["execution_status"]],
        ["機械検証", status["machine_validation"]],
        ["semantic_review_status", status["semantic_review_status"]],
        ["注意", est["note"]],
        [],
        ["構成案", "名称", "estimate_completeness", "区分", "下限", "標準", "上限",
         "算定済", "要確認", "Impossible", "未算定(基準外)", "未採点", "別見積", "対象外", "他へ計上", "既存流用0", "条件付き"],
    ]
    for a in est["architectures"]:
        for label, key, cnt in (("個別ケース", "cases", a["counts"]["cases"]),
                                ("共通作業", "common", a["counts"]["components"]),
                                ("算定済み総計", "computed_total", None)):
            s = a["subtotals"][key]
            line = [a["arch_id"], a["name"], a["estimate_completeness"], label, s["low"], s["std"], s["high"]]
            if cnt is not None:
                line += [cnt.get(k, 0) for k in ("computed", "pending", "impossible", "uncomputed", "unassessed",
                                                 "separate_estimate", "out_of_scope", "counted_elsewhere",
                                                 "reuse_zero", "conditional")]
            rows.append(line)
    conds = conditional_items(assess)
    rows += [[], [f"条件付きの仮定・条件（{len(conds)}件）"], ["区分", "ID", "構成案", "条件", "軸ごとの仮定"]]
    for c in conds:
        rows.append([c["kind"], c["id"], c["arch_ids"], c["conditions"], c["assumptions"]])
    for i, r in enumerate(rows, start=1):
        for j, v in enumerate(r, start=1):
            _cell(ws, i, j, v).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["D"].width = 50
    ws.column_dimensions["E"].width = 50

    # 個別ケース
    ws = wb.create_sheet("個別ケース")
    axis_ids = sorted({k for a in est["architectures"] for c in a["cases"] for k in c.get("axis_scores", {})})
    amap = {(x["case_id"], arch): x for x in assess["case_assessments"] for arch in x["arch_ids"]}
    headers = ["構成案", "case_id", "タイトル", "原項目", "出典"] + \
              [f"{ax} 点" for ax in axis_ids] + ["軸合計", "scope", "estimate_status", "difficulty", "区分",
                                                  "下限", "標準", "上限", "基準参照", "採点理由(軸別)", "総合理由",
                                                  "条件", "確認事項", "備考"]
    data = []
    for a in est["architectures"]:
        for c in a["cases"]:
            case = case_by_id.get(c["case_id"], {})
            refs = [r for sid in case.get("source_item_ids", []) for r in item_by_id.get(sid, {}).get("refs", [])]
            asmt = amap.get((c["case_id"], a["arch_id"]), {})
            axis_reason = "\n".join(
                f"{ax}: {v.get('score')} / 条件:{v.get('condition')} / 根拠:{v.get('evidence')}"
                for ax, v in (asmt.get("axes") or {}).items())
            data.append([a["arch_id"], c["case_id"], c["title"], ", ".join(c["source_item_ids"]), _ref_str(refs)]
                        + [c.get("axis_scores", {}).get(ax) for ax in axis_ids]
                        + [c.get("axis_total"), c.get("scope"), c["estimate_status"], c["difficulty"], c["category"],
                           c["effort"]["low"], c["effort"]["std"], c["effort"]["high"],
                           c.get("standard_refs"), axis_reason, asmt.get("rationale"),
                           asmt.get("conditions"), asmt.get("question_ids"), c.get("notes")])
    _table(ws, headers, data, {3: 36, len(headers) - 4: 60, len(headers) - 3: 40})

    # 共通作業
    ws = wb.create_sheet("共通作業")
    headers = ["構成案", "component_id", "種別", "名称", "責務", "必要性", "実装区分", "変更責務", "難易度",
               "estimate_status", "区分", "基準キー", "単位", "構成要素", "数量", "単価下限", "単価標準", "単価上限",
               "小計下限", "小計標準", "小計上限", "利用ケース", "depends_on", "integrated_into", "cost_owner",
               "判定条件", "理由", "基準参照", "備考"]
    data = []
    for a in est["architectures"]:
        for r in a["components"]:
            comp = comp_by_id[r["component_id"]]
            ca = ca_by_id.get(r["component_id"], {})
            data.append([a["arch_id"], r["component_id"], r["type_id"], comp["name"], comp["responsibility"],
                         comp["need_reason"], r["implementation_kind"], comp.get("change_responsibility"),
                         r["difficulty"], r["estimate_status"], r["category"], r.get("standard_key"), r["unit"],
                         r["elements"], r["quantity"], r["unit_effort"]["low"], r["unit_effort"]["std"],
                         r["unit_effort"]["high"], r["effort"]["low"], r["effort"]["std"], r["effort"]["high"],
                         comp["used_by"], comp["depends_on"], comp["integrated_into"], comp["cost_owner"],
                         ca.get("matched_conditions"), ca.get("rationale"), r.get("standard_refs"), r.get("notes")])
    _table(ws, headers, data, {4: 28, 5: 40})

    # 原項目との対応
    ws = wb.create_sheet("原項目対応")
    cases_by_item: dict[str, list[str]] = {}
    for c in norm["cases"]:
        for sid in c["source_item_ids"]:
            cases_by_item.setdefault(sid, []).append(c["case_id"])
    data = []
    for it in norm["source_items"]:
        data.append(["項目", it["item_id"], it["original_id"], "内部付与" if it["id_assigned"] else "原資料",
                     it["title"], _ref_str(it["refs"]), it["step_count"], it["expected_count"],
                     it["data_pattern_count"], it["disposition"], it["duplicate_of"],
                     cases_by_item.get(it["item_id"]), it["reason"],
                     [f"{x['summary']}（{', '.join(x['question_ids'])}）" for x in it.get("pending_parts", [])]])
    disp = {r["region_id"]: r for r in norm["regions"]}
    for r in inv["regions"]:
        d = disp.get(r["region_id"], {})
        data.append(["領域", r["region_id"], None, None, ", ".join(r["first_row_preview"][:4]),
                     f"{r['workbook_id']}:{r['sheet']}!{r['range']}", None, None, None,
                     f"{d.get('role')}/{d.get('disposition')}" if d else "処遇なし", None, None, d.get("reason"), None])
    _table(ws, ["種別", "ID", "原ID", "ID区分", "タイトル/先頭", "出典", "手順数", "期待結果数", "データパターン数",
                "処遇", "統合先", "対応ケース", "理由", "保留した未確定部分"], data, {5: 40, 6: 30, 13: 40, 14: 40})

    # 画面一覧（PAGEの構成要素とAX-01の操作画面の元）
    screens = screens_doc(rd)
    ws = wb.create_sheet("画面一覧")
    data = [[sc["screen_id"], sc["name"], sc["url"], SCREEN_BASIS_LABELS[sc["basis"]], sc["operated_by"], sc["checked_by"],
             sc["evidence"], _ref_str(sc["refs"])] for sc in screens["screens"]]
    data += [[nc["merged_into"], nc["name"], None, "数えない", None, None, nc["reason"], _ref_str(nc.get("refs", []))]
             for nc in screens["not_counted"]]
    _table(ws, ["screen_id", "画面", "URL", "根拠", "操作するケース", "確認だけのケース", "判断の根拠", "出典"], data,
           {2: 28, 7: 50, 8: 30})

    # 確認事項
    ws = wb.create_sheet("確認事項")
    _table(ws, ["question_id", "種類", "対象", "内容", "影響", "仮の扱い", "状態"],
           [[q["question_id"], q["kind"], q["targets"], q["text"], q.get("impact"), q.get("proposed_default"),
             q["status"]] for q in questions], {4: 60, 5: 40, 6: 40})

    # レビュー（人の確認用）
    ws = wb.create_sheet("レビュー")
    data = []
    for c in norm["cases"]:
        src = "\n".join(it["title"] for sid in c["source_item_ids"] if (it := item_by_id.get(sid)))
        for a in est["architectures"]:
            row = next((x for x in a["cases"] if x["case_id"] == c["case_id"]), None)
            asmt = amap.get((c["case_id"], a["arch_id"]), {})
            data.append([a["arch_id"], c["case_id"], src, c["relation"], c.get("relation_reason"),
                         "\n".join(c["steps"]), "\n".join(c["expected"]),
                         row.get("axis_scores") if row else None, row["difficulty"] if row else None,
                         asmt.get("rationale"), None, None, None])
    _table(ws, ["構成案", "case_id", "原項目", "分解/統合", "理由", "手順", "期待結果", "5軸", "難易度", "AI理由",
                "レビュー判定(OK/NG)", "正しい値", "コメント"], data, {3: 30, 6: 40, 7: 40, 10: 40, 13: 40})

    # 意味レビューの指摘と対応
    findings = review_findings(rd)
    ws = wb.create_sheet("意味レビュー")
    _table(ws, ["対象", "関連", "種類", "指摘", "修正案", "対応", "対応内容", "質問"],
           [[f["target"], f.get("related"), f["kind"], f["message"], f["suggestion"],
             RESOLUTION_LABELS.get((f.get("resolution") or {}).get("status"), "未記録"),
             (f.get("resolution") or {}).get("note"), (f.get("resolution") or {}).get("question_ids")]
            for f in findings], {4: 60, 5: 40, 7: 40})

    # 検証結果
    ws = wb.create_sheet("検証結果")
    data = []
    for stage, v in val["machine"]["stages"].items():
        for sev in ("errors", "warnings"):
            for e in v[sev]:
                data.append([stage, e.get("severity", sev[:-1]), e["code"], e["path"], e["message"]])
    _table(ws, ["工程", "重大度", "コード", "位置", "内容"], data, {4: 30, 5: 80})

    # 実行条件
    ws = wb.create_sheet("実行条件")
    rows = [["項目", "値"], ["run_id", man["run_id"]], ["created_at", man["created_at"]]]
    for i in man["inputs"]:
        rows.append([f"入力 {i['workbook_id']}", f"{i['path']} sha256={i['sha256']} size={i['size']}"])
    for k in ("cases", "common", "adapter"):
        s = man["standards"][k]
        rows.append([f"基準 {k}", f"{s['path']} sha256={s['sha256']} 版={s.get('version')}"])
    rows.append(["設定", f"{man['config']['path']} sha256={man['config']['sha256']}"])
    for k, v in man["harness"].items():
        rows.append([f"harness.{k}", v])
    for k, v in man["environment"].items():
        rows.append([f"env.{k}", v])
    rows.append(["rule_versions", est["rule_versions"]])
    for i, r in enumerate(rows, start=1):
        for j, v in enumerate(r, start=1):
            _cell(ws, i, j, v)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 120

    atomic_save_workbook(wb, rd / "estimate.xlsx")
    annotated = export_annotated(rd, est, assess, norm, questions)
    write_text(rd / "summary.md", render_summary(man, est, val, questions, conds, norm, findings, screens))
    return {"xlsx": str(rd / "estimate.xlsx"), "summary": str(rd / "summary.md"), "annotated": annotated}


DIFFICULTY_LABELS = {"Low": "低", "Mid": "中", "High": "高", "Impossible": "不可能"}
DIFFICULTY_ORDER = ["Low", "Mid", "High", "Impossible"]
ANNOTATE_HEADERS = ["自動化難易度", "自動化工数(標準)", "懸念点", "備考"]


def _fmt_num(v) -> str:
    return f"{v:g}" if isinstance(v, (int, float)) else str(v)


def _item_annotation(item: dict, cases: list[dict], rows_by_arch: dict, amap: dict, qtext: dict,
                     owner_of_case: dict, arch_names: dict, item_by_id: dict) -> dict[str, Any]:
    """原項目1件ぶんの追記内容。値はすべて estimate.json / ai/*.json から転記し、ここで計算しない（合算のみ）。

    備考・懸念点は内部ID（構成案ID・ケースID・質問ID）を出さず、Excelだけで読める日本語にする。
    構成案が1つのときは構成案名も出さない。
    """
    def label(item_id: str) -> str:
        it = item_by_id[item_id]
        return f"{it['original_id'] or item_id}「{it['title']}」" if it["title"] else (it["original_id"] or item_id)

    if not cases:
        return {"difficulty": {}, "effort": {}, "concerns": [], "notes": [f"見積もり対象外（処遇: {item['disposition']}）"]
                + ([item["reason"]] if item.get("reason") else [])}
    difficulty: dict[str, str | None] = {}
    effort: dict[str, float | None] = {}
    concerns: list[str] = []
    notes: list[str] = []
    for arch_id, rows in rows_by_arch.items():
        pre = f"［{arch_names[arch_id]}］" if arch_names[arch_id] else ""
        mine = [rows[c["case_id"]] for c in cases if c["case_id"] in rows]
        diffs = [r["difficulty"] for r in mine if r.get("difficulty")]
        difficulty[arch_id] = max(diffs, key=DIFFICULTY_ORDER.index) if diffs else None
        owned = [r for r in mine if owner_of_case[r["case_id"]] == item["item_id"]]
        counted = [r for r in owned if r["effort"]["std"] is not None]
        effort[arch_id] = sum(r["effort"]["std"] for r in counted) if counted else None
        if counted:
            lo = sum(r["effort"]["low"] for r in counted)
            hi = sum(r["effort"]["high"] for r in counted)
            notes.append(f"{pre}工数の幅: 下限{_fmt_num(lo)}〜上限{_fmt_num(hi)}人時")
        for r in mine:
            if r["estimate_status"] == "条件付き":
                notes.append(f"{pre}条件付きの見積もり（懸念点の仮定が成り立つ場合の数字）")
            elif r["estimate_status"] == "要確認":
                notes.append(f"{pre}要確認のため難易度・工数は未算定（懸念点の確認事項が決まるまで算定しない）")
            if r.get("scope") not in (None, "対象内"):
                notes.append(f"{pre}今回の見積もりの{r['scope']}")
            if r.get("notes"):
                notes.append(f"{pre}{r['notes']}")
            owner = owner_of_case[r["case_id"]]
            if r["effort"]["std"] is not None and owner != item["item_id"]:
                notes.append(f"{pre}{label(owner)}などと合わせて1つの見積もりにまとめたため、工数は{label(owner)}の行に記載")
            a = amap.get((r["case_id"], arch_id), {})
            concerns += [f"{pre}{c}" for c in a.get("conditions", [])]
            concerns += [f"{pre}{ax}の仮定: {x}" for ax, v in (a.get("axes") or {}).items()
                         for x in v.get("assumptions", [])]
            concerns += [f"{pre}要確認: {qtext.get(q, q)}" for q in a.get("question_ids", [])]
    for c in cases:
        others = [i for i in c["source_item_ids"] if i != item["item_id"]]
        if others and item["item_id"] == owner_of_case[c["case_id"]]:
            notes.append("次の項目と合わせて1つの見積もりにまとめ、その工数をこの行に記載: "
                         + "、".join(label(i) for i in others))
    n_own = [c for c in cases if c["source_item_ids"] == [item["item_id"]]]
    if len(n_own) > 1:
        notes.append(f"この項目は{len(n_own)}つの見積もり単位に分解して見積もり（工数は合算）")
    return {"difficulty": difficulty, "effort": effort,
            "concerns": list(dict.fromkeys(concerns)), "notes": list(dict.fromkeys(notes))}


def export_annotated(rd: Path, est: dict, assess: dict, norm: dict, questions: list[dict]) -> list[str]:
    """元の項目書の各項目の右隣に、難易度・工数・懸念点・備考の列を足したコピーを annotated/ へ出す。

    元ファイルは変更しない（コピーに追記）。共通作業の工数は項目に紐づかないため含めない。
    """
    man = read_json(rd / "manifest.json")
    case_by_id = {c["case_id"]: c for c in norm["cases"]}
    amap = {(x["case_id"], arch): x for x in assess["case_assessments"] for arch in x["arch_ids"]}
    qtext = {q["question_id"]: q["text"] for q in questions}
    rows_by_arch = {a["arch_id"]: {c["case_id"]: c for c in a["cases"]} for a in est["architectures"]}
    multi = len(rows_by_arch) > 1
    arch_names = {a["arch_id"]: (a["name"] if multi else "") for a in est["architectures"]}
    item_by_id = {i["item_id"]: i for i in norm["source_items"]}
    cases_by_item: dict[str, list[dict]] = {}
    for c in norm["cases"]:
        for sid in c["source_item_ids"]:
            cases_by_item.setdefault(sid, []).append(c)
    # 統合ケース（複数項目→1ケース）の工数は先頭の項目にだけ計上し、二重に数えない
    owner_of_case = {c["case_id"]: c["source_item_ids"][0] for c in norm["cases"] if c["source_item_ids"]}
    header_rows = {r["region_id"]: r.get("header_rows") or [] for r in norm["regions"]}
    outputs: list[str] = []
    for inp in man["inputs"]:
        src = Path(inp["path"])
        if not src.exists() or src.suffix.lower() not in (".xlsx", ".xlsm"):
            continue
        wb = load_workbook(src)  # マクロは保持せず(.xlsx出力)、実行もしない
        placed: dict[str, int] = {}  # シート名 -> 追記開始列
        headers_done: set[str] = set()
        for it in norm["source_items"]:
            ref = next((r for r in it["refs"] if r["workbook_id"] == inp["workbook_id"]), None)
            if ref is None or ref["sheet"] not in wb.sheetnames:
                continue
            ws = wb[ref["sheet"]]
            col0 = placed.setdefault(ref["sheet"], ws.max_column + 1)
            first_row = int(re.search(r"\d+", ref["range"].split(":")[0]).group())
            ann = _item_annotation(it, cases_by_item.get(it["item_id"], []), rows_by_arch, amap, qtext, owner_of_case,
                                  arch_names, item_by_id)
            archs = list(rows_by_arch)
            diff = "\n".join((f"［{arch_names[a]}］" if multi else "") + (DIFFICULTY_LABELS.get(ann["difficulty"].get(a)) or "判定不可")
                             for a in archs) if ann["difficulty"] else None
            eff = "\n".join((f"［{arch_names[a]}］" if multi else "") + ("" if ann["effort"].get(a) is None else _fmt_num(ann["effort"][a]))
                            for a in archs) if ann["effort"] else None
            if not multi and ann["effort"]:
                eff = ann["effort"][archs[0]]
            values = [diff, eff, ann["concerns"], ann["notes"]]
            for j, v in enumerate(values):
                c = _cell(ws, first_row, col0 + j, v)
                c.alignment = Alignment(wrap_text=True, vertical="top")
            hdr = next((max(header_rows.get(rid, []) or [0]) for rid in it["region_ids"] if header_rows.get(rid)), 0)
            if hdr and ref["sheet"] not in headers_done and hdr < first_row:
                headers_done.add(ref["sheet"])
                for j, h in enumerate(ANNOTATE_HEADERS):
                    c = _cell(ws, hdr, col0 + j, h)
                    c.font = Font(bold=True)
                    c.fill = HEADER_FILL
                    c.alignment = Alignment(wrap_text=True, vertical="top")
        for sheet, col0 in placed.items():
            for j, w in enumerate((14, 16, 40, 40)):
                wb[sheet].column_dimensions[get_column_letter(col0 + j)].width = w
        out = rd / "annotated" / f"{src.stem}_見積付き.xlsx"
        atomic_save_workbook(wb, out)
        outputs.append(str(out))
    return outputs


SCREEN_BASIS_LABELS = {"inferred": "項目書から推定", "observed": "サイト調査で確認"}


def screens_doc(rd: Path) -> dict:
    p = rd / "ai/screens.json"
    return read_json(p) if p.exists() else {"screens": [], "not_counted": []}


def _screen_lines(screens: dict) -> list[str]:
    if not screens["screens"]:
        return []
    basis: dict[str, int] = {}
    for sc in screens["screens"]:
        basis[sc["basis"]] = basis.get(sc["basis"], 0) + 1
    return [f"- 画面一覧: {len(screens['screens'])}画面（"
            + "、".join(f"{SCREEN_BASIS_LABELS[k]} {v}" for k, v in sorted(basis.items()))
            + f"）、数えなかった候補 {len(screens['not_counted'])}件"]


RESOLUTION_LABELS = {"fixed": "修正済", "question": "質問化", "rejected": "不採用", "deferred": "未対応"}


def review_findings(rd: Path) -> list[dict]:
    p = rd / "ai/semantic_review.json"
    return read_json(p).get("findings", []) if p.exists() else []


def _review_lines(findings: list[dict]) -> list[str]:
    """意味レビューの指摘と対応。未対応（deferred・未記録）を先に出す。"""
    if not findings:
        return []
    label = lambda f: RESOLUTION_LABELS.get((f.get("resolution") or {}).get("status"), "未記録")  # noqa: E731
    open_f = [f for f in findings if label(f) in ("未対応", "未記録")]
    out = [f"## 意味レビューの指摘 ({len(findings)}件、うち未対応{len(open_f)}件)", ""]
    for f in sorted(findings, key=lambda f: label(f) not in ("未対応", "未記録")):
        note = (f.get("resolution") or {}).get("note")
        out.append(f"- [{label(f)}] {f['kind']} {f['target']}: {f['message'][:120]}" + (f" → {note[:80]}" if note else ""))
    return out + [""]


def _fmt(v):
    return "—" if v is None else (f"{v:g}" if isinstance(v, float) else str(v))


def _warning_lines(val: dict) -> list[str]:
    from .validate import summarize_warnings

    ws = summarize_warnings(val["machine"])
    if not ws:
        return []
    return ["- 警告の内訳（件数。詳細は estimate.xlsx の「検証結果」）:"] + \
           [f"  - {w['stage']} / {w['kind']}: {w['count']}件（例 {w['sample'][:60]}）" for w in ws]


def _mapping_lines(norm: dict | None) -> list[str]:
    """原項目と見積もりケースの対応件数。分解・統合で件数が変わったことを読み手に示す。"""
    if not norm:
        return []
    items = norm.get("source_items", [])
    cases = norm.get("cases", [])
    per_item: dict[str, int] = {}
    for c in cases:
        for sid in c.get("source_item_ids", []):
            per_item[sid] = per_item.get(sid, 0) + 1
    split = sorted(i for i, n in per_item.items() if n > 1)
    merged = sorted(c["case_id"] for c in cases if len(c.get("source_item_ids", [])) > 1)
    line = f"- 原項目 {len(items)}件 → 見積もりケース {len(cases)}件"
    if split:
        line += f"（分解された原項目: {', '.join(split)}）"
    if merged:
        line += f"（統合されたケース: {', '.join(merged)}）"
    pending = [(i["item_id"], x) for i in items for x in i.get("pending_parts", [])]
    out = [line]
    if pending:
        out.append(f"- ケースにせず保留した未確定部分 {len(pending)}件（未解決件数に含む）: "
                   + "; ".join(f"{iid} {x['summary'][:40]}（{', '.join(x['question_ids'])}）" for iid, x in pending))
    return out


def render_summary(man: dict, est: dict, val: dict, questions: list[dict], conds: list[dict],
                   norm: dict | None = None, findings: list[dict] | None = None, screens: dict | None = None) -> str:
    st = val["status"]
    out = [f"# 見積もり結果 {man['run_id']}", "",
           f"- 入力: {', '.join(i['file_name'] for i in man['inputs'])}",
           f"- execution_status: {st['execution_status']} / 機械検証: {st['machine_validation']} "
           f"(エラー{val['machine']['error_count']}件, 警告{val['machine']['warning_count']}件)",
           f"- semantic_review_status: {st['semantic_review_status']}",
           *_mapping_lines(norm),
           *_screen_lines(screens or {"screens": [], "not_counted": []}),
           *_warning_lines(val),
           f"- 基準: 個別 {man['standards']['cases'].get('version')} / 共通 {man['standards']['common'].get('version')}",
           "", f"> {est['note']}", ""]
    for a in est["architectures"]:
        out += [f"## 構成案 {a['arch_id']}: {a['name']} — {a['estimate_completeness']}", "",
                "| 区分 | 下限 | 標準 | 上限 |", "|---|---:|---:|---:|"]
        for label, key in (("個別ケース", "cases"), ("共通作業", "common"), ("算定済み総計", "computed_total")):
            s = a["subtotals"][key]
            out.append(f"| {label} | {_fmt(s['low'])} | {_fmt(s['std'])} | {_fmt(s['high'])} |")
        cc, mc = a["counts"]["cases"], a["counts"]["components"]
        fmt = lambda d: ", ".join(f"{k}={v}" for k, v in d.items())  # noqa: E731
        out += ["", f"- 個別ケース件数: {fmt(cc)}", f"- 共通作業件数: {fmt(mc)}", f"- 未解決件数: {a['counts']['unresolved_items']}", ""]
    out += [f"## 条件付きの仮定・条件 ({len(conds)}件)", ""]
    for c in conds:
        detail = "; ".join(c["conditions"] + c["assumptions"]) or "（条件の記載なし）"
        out.append(f"- [{c['kind']}] {c['id']}: {detail}")
    out.append("")
    out += _review_lines(findings or [])
    open_q = [q for q in questions if q["status"] != "resolved"]
    out += [f"## 確認事項・仮定 ({len(open_q)}件未解決)", ""]
    for q in open_q[:50]:
        out.append(f"- [{q['kind']}] {q['question_id']}: {q['text']}")
    return "\n".join(out) + "\n"
