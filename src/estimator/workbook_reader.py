"""Excel(.xlsx)の構造抽出。

意味判断は一切行わず、セル・数式・保存済み計算結果・結合・非表示・コメント・図形の位置を
原文のまま保持する。顧客項目書・基準Excelの双方の下位処理として使う。
"""

from __future__ import annotations

import datetime as _dt
import posixpath
import re
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries

from .io_utils import sha256_file

SUPPORTED_SUFFIXES = {".xlsx", ".xlsm"}  # .xlsm はOOXMLとして読む。VBAは実行しない

_NS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
}
_OLE_MAGIC = bytes.fromhex("D0CF11E0A1B11AE1")


class UnsupportedWorkbook(Exception):
    """MVPで正式対応しない形式（.xls、暗号化、破損など）。"""


def _json_value(v: Any) -> Any:
    if isinstance(v, (_dt.datetime, _dt.date, _dt.time)):
        return v.isoformat()
    if isinstance(v, _dt.timedelta):
        return v.total_seconds()
    if isinstance(v, (str, int, float, bool)) or v is None:
        return v
    return str(v)


def check_supported(path: Path) -> None:
    if not path.exists():
        raise FileNotFoundError(str(path))
    with open(path, "rb") as f:
        head = f.read(8)
    if head == _OLE_MAGIC:
        raise UnsupportedWorkbook(
            "OLE形式です（.xls、またはパスワード保護により暗号化された.xlsx）。MVPでは非対応です。"
        )
    if path.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise UnsupportedWorkbook(f"拡張子 {path.suffix} はMVPの正式対応外です（.xlsx / .xlsm のみ）。")
    if not zipfile.is_zipfile(path):
        raise UnsupportedWorkbook("ZIP(OOXML)として読めません。破損または別形式の可能性があります。")


def _drawings_by_sheet(path: Path) -> dict[str, list[dict]]:
    """図・画像・グラフ・図形のアンカー位置をOOXMLから直接取得する（Pillow等に依存しない）。"""
    result: dict[str, list[dict]] = {}
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        try:
            wb_xml = ET.fromstring(z.read("xl/workbook.xml"))
            wb_rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        except KeyError:
            return result
        rid_to_target = {r.get("Id"): r.get("Target") for r in wb_rels.findall("rel:Relationship", _NS)}
        for sh in wb_xml.findall("main:sheets/main:sheet", _NS):
            sname = sh.get("name")
            target = rid_to_target.get(sh.get(f"{{{_NS['r']}}}id"), "")
            sheet_path = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
            rels_path = posixpath.join(posixpath.dirname(sheet_path), "_rels", posixpath.basename(sheet_path) + ".rels")
            items: list[dict] = []
            if rels_path in names:
                srels = ET.fromstring(z.read(rels_path))
                for rel in srels.findall("rel:Relationship", _NS):
                    rtype = rel.get("Type", "").rsplit("/", 1)[-1]
                    if rtype not in ("drawing", "vmlDrawing", "oleObject", "package", "image"):
                        continue
                    tgt = rel.get("Target", "")
                    dpath = tgt.lstrip("/") if tgt.startswith("/") else posixpath.normpath(posixpath.join(posixpath.dirname(sheet_path), tgt))
                    if rtype != "drawing":
                        if rtype in ("oleObject", "package"):
                            items.append({"kind": "embedded_object", "part": dpath, "anchor": None})
                        continue
                    if dpath not in names:
                        continue
                    dxml = ET.fromstring(z.read(dpath))
                    for anchor in list(dxml):
                        tag = anchor.tag.split("}")[-1]
                        if tag not in ("twoCellAnchor", "oneCellAnchor", "absoluteAnchor"):
                            continue
                        kind = "shape"
                        for child in anchor:
                            ctag = child.tag.split("}")[-1]
                            if ctag == "pic":
                                kind = "image"
                            elif ctag == "graphicFrame":
                                kind = "chart_or_graphic"
                            elif ctag == "grpSp":
                                kind = "group"
                            elif ctag == "sp":
                                kind = "shape"
                        text = " ".join(t.text for t in anchor.iter() if t.tag.endswith("}t") and t.text)
                        frm = anchor.find("xdr:from", _NS)
                        ref = None
                        if frm is not None:
                            col = int(frm.findtext("xdr:col", "0", _NS)) + 1
                            row = int(frm.findtext("xdr:row", "0", _NS)) + 1
                            ref = f"{get_column_letter(col)}{row}"
                        items.append({"kind": kind, "anchor": ref, "text": text or None, "part": dpath})
            result[sname] = items
    return result


def _hidden_cols(ws) -> list[str]:
    cols: set[int] = set()
    for key, dim in ws.column_dimensions.items():
        if dim.hidden:
            lo = dim.min or openpyxl.utils.column_index_from_string(key)
            hi = dim.max or lo
            cols.update(range(lo, hi + 1))
    return [get_column_letter(c) for c in sorted(cols)]


def read_workbook(path: str | Path, workbook_id: str = "W1") -> dict:
    """ブックを読み、原文・構造を保持したdictを返す。数式は評価しない。"""
    path = Path(path)
    check_supported(path)
    warnings: list[dict] = []
    try:
        wb_f = openpyxl.load_workbook(path, data_only=False, rich_text=False)
        wb_v = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:  # noqa: BLE001 - 原因をそのまま制約として報告する
        raise UnsupportedWorkbook(f"openpyxlで読めません: {type(e).__name__}: {e}") from e

    drawings = _drawings_by_sheet(path)
    with zipfile.ZipFile(path) as z:
        vba_parts = [n for n in z.namelist() if n.lower().endswith("vbaproject.bin")]
    if vba_parts:
        warnings.append({"code": "VBA_PRESENT", "ref": ",".join(vba_parts),
                         "message": "VBAマクロを含みます。マクロは実行・解析せず、セルの値だけを読みます"})
    sheets = []
    for idx, ws in enumerate(wb_f.worksheets, start=1):
        wsv = wb_v[ws.title]
        merged = []
        merged_cover: dict[str, str] = {}
        for mr in ws.merged_cells.ranges:
            min_col, min_row, max_col, max_row = range_boundaries(str(mr))
            anchor = f"{get_column_letter(min_col)}{min_row}"
            merged.append({"range": str(mr), "anchor": anchor})
            for r in range(min_row, max_row + 1):
                for c in range(min_col, max_col + 1):
                    ref = f"{get_column_letter(c)}{r}"
                    if ref != anchor:
                        merged_cover[ref] = str(mr)
        cells = []
        comments = []
        for row in ws.iter_rows():
            for cell in row:
                if cell.comment is not None:
                    comments.append({"ref": cell.coordinate, "text": cell.comment.text, "author": cell.comment.author})
                v = cell.value
                if v is None or (isinstance(v, str) and v == ""):
                    continue
                entry: dict[str, Any] = {"ref": cell.coordinate, "row": cell.row, "col": cell.column}
                if cell.data_type == "f" or (isinstance(v, str) and v.startswith("=") and cell.data_type == "f"):
                    cached = wsv[cell.coordinate].value
                    entry.update(
                        type="formula",
                        formula=str(v),
                        cached_value=_json_value(cached),
                        cached_missing=cached is None,
                    )
                    if cached is None:
                        warnings.append({"code": "FORMULA_NO_CACHE", "ref": f"{ws.title}!{cell.coordinate}",
                                         "message": "数式の保存済み計算結果がありません（0扱いしません）"})
                elif hasattr(v, "text") and not isinstance(v, str):  # 配列数式等
                    entry.update(type="formula", formula=str(getattr(v, "text", v)), cached_value=None, cached_missing=True)
                else:
                    entry["type"] = {"n": "number", "s": "string", "b": "bool", "d": "date", "e": "error"}.get(cell.data_type, cell.data_type)
                    entry["value"] = _json_value(v)
                    if isinstance(v, str) and "\n" in v:
                        entry["has_newline"] = True
                cells.append(entry)
        hidden_rows = sorted(r for r, d in ws.row_dimensions.items() if d.hidden)
        dv = []
        try:
            for d in ws.data_validations.dataValidation:
                dv.append({"sqref": str(d.sqref), "type": d.type, "formula1": d.formula1})
        except AttributeError:
            pass
        sheets.append({
            "index": idx,
            "name": ws.title,
            "state": ws.sheet_state,
            "max_row": ws.max_row,
            "max_col": ws.max_column,
            "hidden_rows": hidden_rows,
            "hidden_cols": _hidden_cols(ws),
            "merged": merged,
            "merged_cover": merged_cover,
            "cells": cells,
            "comments": comments,
            "drawings": drawings.get(ws.title, []),
            "data_validations": dv,
        })
        if ws.sheet_state != "visible":
            warnings.append({"code": "HIDDEN_SHEET", "ref": ws.title, "message": f"シート状態={ws.sheet_state}"})
        if hidden_rows or sheets[-1]["hidden_cols"]:
            warnings.append({"code": "HIDDEN_ROWS_COLS", "ref": ws.title,
                             "message": f"非表示行={hidden_rows} 非表示列={sheets[-1]['hidden_cols']}"})
        for d in sheets[-1]["drawings"]:
            warnings.append({"code": "DRAWING_PRESENT", "ref": f"{ws.title}!{d.get('anchor')}",
                             "message": f"{d['kind']}があります。内容はテキスト化できない場合があります"})
    chartsheets = [cs.title for cs in getattr(wb_f, "chartsheets", [])]
    for cs in chartsheets:
        warnings.append({"code": "CHARTSHEET", "ref": cs, "message": "グラフシートは解析対象外です"})
    return {
        "workbook_id": workbook_id,
        "path": str(path),
        "file_name": path.name,
        "sha256": sha256_file(path),
        "size": path.stat().st_size,
        "macro_enabled_format": path.suffix.lower() == ".xlsm",
        "vba_present": bool(vba_parts),
        "sheets": sheets,
        "chartsheets": chartsheets,
        "warnings": warnings,
    }


def cell_text(entry: dict) -> str:
    if entry["type"] == "formula":
        cv = entry.get("cached_value")
        return f"{entry['formula']} [保存値: {'なし' if entry.get('cached_missing') else cv}]"
    v = entry.get("value")
    return "" if v is None else str(v)


def detect_regions(sheet: dict, workbook_id: str) -> list[dict]:
    """連続する非空行のまとまりを領域候補とする（意味判断はしない）。"""
    rows_used: dict[int, list[dict]] = {}
    for c in sheet["cells"]:
        rows_used.setdefault(c["row"], []).append(c)
    occupied = set(rows_used)
    for m in sheet["merged"]:
        min_col, min_row, max_col, max_row = range_boundaries(m["range"])
        if any(r in rows_used for r in range(min_row, max_row + 1)):
            occupied.update(range(min_row, max_row + 1))
    blocks: list[list[int]] = []
    for r in sorted(occupied):
        if blocks and r == blocks[-1][-1] + 1:
            blocks[-1].append(r)
        else:
            blocks.append([r])
    regions = []
    for i, rows in enumerate(blocks, start=1):
        cells = [c for r in rows for c in rows_used.get(r, [])]
        min_c = min((c["col"] for c in cells), default=1)
        max_c = max((c["col"] for c in cells), default=1)
        first_row_cells = sorted(rows_used.get(rows[0], []), key=lambda c: c["col"])
        rng = f"{get_column_letter(min_c)}{rows[0]}:{get_column_letter(max_c)}{rows[-1]}"
        regions.append({
            "region_id": f"{workbook_id}-S{sheet['index']}-R{i:03d}",
            "workbook_id": workbook_id,
            "sheet": sheet["name"],
            "range": rng,
            "first_row": rows[0],
            "last_row": rows[-1],
            "row_count": len(rows),
            "nonempty_cells": len(cells),
            "hidden_rows": [r for r in rows if r in sheet["hidden_rows"]],
            "merged_ranges": [m["range"] for m in sheet["merged"]
                              if rows[0] <= range_boundaries(m["range"])[1] <= rows[-1]],
            "first_row_preview": [cell_text(c)[:40] for c in first_row_cells[:12]],
        })
    for j, d in enumerate(sheet["drawings"], start=1):
        regions.append({
            "region_id": f"{workbook_id}-S{sheet['index']}-D{j:03d}",
            "workbook_id": workbook_id,
            "sheet": sheet["name"],
            "range": d.get("anchor"),
            "kind": "drawing",
            "drawing_kind": d["kind"],
            "drawing_text": d.get("text"),
            "first_row": None, "last_row": None, "row_count": 0, "nonempty_cells": 0,
            "hidden_rows": [], "merged_ranges": [], "first_row_preview": [],
        })
    return regions


def render_sheet(sheet: dict, row_from: int | None = None, row_to: int | None = None,
                 max_len: int | None = None) -> str:
    """AIが読むためのテキスト表現。座標・結合・非表示を明示し、値は改変しない（改行は⏎）。"""
    lines = [f"## シート: {sheet['name']} (index={sheet['index']}, state={sheet['state']}, "
             f"max_row={sheet['max_row']}, max_col={sheet['max_col']})"]
    if sheet["hidden_cols"]:
        lines.append(f"非表示列: {', '.join(sheet['hidden_cols'])}")
    rows: dict[int, list[dict]] = {}
    for c in sheet["cells"]:
        rows.setdefault(c["row"], []).append(c)
    merged_by_anchor = {m["anchor"]: m["range"] for m in sheet["merged"]}
    lo = row_from or 1
    hi = row_to or sheet["max_row"]
    hidden = set(sheet["hidden_rows"])
    for r in range(lo, hi + 1):
        parts = []
        for c in sorted(rows.get(r, []), key=lambda x: x["col"]):
            t = cell_text(c).replace("\r\n", "\n").replace("\n", "⏎")
            if max_len and len(t) > max_len:
                t = t[:max_len] + "…"
            tag = f"[結合 {merged_by_anchor[c['ref']]}]" if c["ref"] in merged_by_anchor else ""
            parts.append(f"{c['ref']}{tag}={t!r}")
        covered = sorted({rng for ref, rng in sheet["merged_cover"].items()
                          if int(re.sub(r"[A-Z]+", "", ref)) == r})
        if not parts and not covered:
            continue
        flag = " [非表示行]" if r in hidden else ""
        cov = f" (結合範囲内: {', '.join(covered)})" if covered else ""
        lines.append(f"R{r}{flag}: " + " | ".join(parts) + cov)
    for cm in sheet["comments"]:
        if lo <= int(re.sub(r"[A-Z]+", "", cm["ref"])) <= hi:
            lines.append(f"コメント {cm['ref']}: {cm['text']!r}")
    for d in sheet["drawings"]:
        lines.append(f"図形 {d['kind']} @ {d.get('anchor')}: text={d.get('text')!r}")
    return "\n".join(lines)
